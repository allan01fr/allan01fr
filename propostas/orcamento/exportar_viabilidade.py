"""
Planilha do estudo de viabilidade (padrão KAPTEK, fórmulas vivas).

Abas: Resumo (indicadores por alternativa), Premissas, Alternativas, FluxoCaixa (anos 0 a 30,
limitados pelo horizonte) e Sensibilidade. Mudar tarifa, TMA, horizonte, reajustes, CAPEX ou a
redução de consumo recalcula tudo, inclusive payback, breakeven, VPL, TIR e LCC.
Funções compatíveis com o Excel 2007 (SUMPRODUCT, COUNTIF, INDEX, IRR, IFERROR).
"""

import io

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.utils import get_column_letter as L

from .exportar_xlsx import AZUL, CINZA, FUNDO_CAB, INT, MOEDA, NUM, PCT, Folha, _impressao, _logo, _org
from .viabilidade import HORIZONTE_MAX, VARIACOES_CAPEX, VARIACOES_ECONOMIA, VARIACOES_TARIFA, completar

KWH = '#,##0;(#,##0);"-"'
ANOS = '0.00 "anos";-0.00;"-"'


def gerar(proposta, estudo):
    """Bytes do .xlsx do estudo. `estudo` = viabilidade.calcular(...): dá os valores de partida e o CAPEX já
    resolvido quando ele está vinculado ao preço de uma opção da proposta."""
    v = completar(proposta.get("viabilidade"))
    b, fin = v["base"], v["financeiro"]
    alts = estudo["alternativas"]
    wb = Workbook()
    ws_res = wb.active
    ws_res.title = "Resumo"

    # --------------------------------------------------------------- Premissas
    ws = wb.create_sheet("Premissas")
    f = Folha(ws)
    f.larguras({"A": 4, "B": 52, "C": 18, "D": 60})
    f.titulo(1, "PREMISSAS DO ESTUDO DE VIABILIDADE", 4)
    ref = {}
    linhas = [
        ("sec", "Linha de base — sistema atual"),
        ("modo", "Modo da linha de base (tr | kwh)", b["modo"], None, "tr = TR × kW/TR × horas × fator de carga; kwh = consumo anual informado"),
        ("tr", "Capacidade instalada (TR)", float(b["tr"]), NUM, ""),
        ("kw_real", "Eficiência medida (kW/TR)", float(b["kw_tr_real"]), NUM, "Desempenho real do sistema"),
        ("kw_proj", "Eficiência de projeto (kW/TR)", float(b["kw_tr_projeto"]), NUM, "Estimativa do projeto original"),
        ("horas", "Horas de operação por ano", float(b["horas_ano"]), KWH, ""),
        ("fc", "Fator de carga médio", float(b["fator_carga"]), NUM, ""),
        ("kwh_inf", "Consumo anual informado (kWh)", float(b["kwh_ano"]), KWH, "Usado no modo kwh"),
        ("tarifa", "Tarifa média de energia (R$/kWh)", float(b["tarifa"]), NUM, "Fonte: ANEEL / faturas do cliente"),
        ("manut", "Custo anual de manutenção atual (R$)", float(b["manutencao_ano"]), MOEDA, ""),
        ("paliativos", "Paliativos / manutenção de emergência (R$)", float(b.get("paliativos") or 0), MOEDA, "Comuns a todas as alternativas: entram no LCC"),
        ("sec", "Financeiro"),
        ("tma", "Taxa mínima de atratividade (TMA, a.a.)", float(fin["tma"]), PCT, "Custo de capital do cliente"),
        ("h", f"Horizonte de análise (anos, até {HORIZONTE_MAX})", int(fin["horizonte"]), INT, "Vida útil considerada"),
        ("ge", "Reajuste anual da energia", float(fin["reajuste_energia"]), PCT, ""),
        ("gm", "Reajuste anual da manutenção", float(fin["reajuste_manutencao"]), PCT, ""),
    ]
    r = 3
    for item in linhas:
        if item[0] == "sec":
            f.secao(r, item[1], 4)
            r += 1
            continue
        chave, rot, val, fmt, obs = item
        f.c(r, 2, rot)
        f.entrada(r, 3, val, fmt)
        f.c(r, 4, obs, cor=CINZA, tam=8)
        ref[chave] = f"Premissas!$C${r}"
        r += 1
    r += 1
    f.secao(r, "Resultados da linha de base", 4)
    r += 1
    derivados = [
        ("cb", "Consumo atual (kWh/ano)", f"=IF({ref['modo']}=\"tr\",{ref['tr']}*{ref['kw_real']}*{ref['horas']}*{ref['fc']},{ref['kwh_inf']})", KWH),
        ("cp", "Consumo previsto no projeto original (kWh/ano)", f"=IF({ref['modo']}=\"tr\",{ref['tr']}*{ref['kw_proj']}*{ref['horas']}*{ref['fc']},0)", KWH),
        ("excesso", "Consumo acima do projeto", None, PCT),
        ("perda", "Custo anual da perda de desempenho (R$)", None, MOEDA),
        ("energia_base", "Custo anual de energia atual (R$)", None, MOEDA),
    ]
    for chave, rot, formula, fmt in derivados:
        ref[chave] = f"Premissas!$C${r}"
        r += 1
    calc = {
        "excesso": f"=IF({ref['cp']}>0,{ref['cb']}/{ref['cp']}-1,0)",
        "perda": f"=IF({ref['cp']}>0,({ref['cb']}-{ref['cp']})*{ref['tarifa']},0)",
        "energia_base": f"={ref['cb']}*{ref['tarifa']}",
    }
    for chave, rot, formula, fmt in derivados:
        rr = int(ref[chave].split("$")[-1])
        f.c(rr, 2, rot, negrito=True)
        f.c(rr, 3, formula or calc[chave], fmt, negrito=True, fundo=FUNDO_CAB)
    ws.freeze_panes = "A3"
    _impressao(ws, paisagem=False)

    # --------------------------------------------------------------- Alternativas
    ws = wb.create_sheet("Alternativas")
    f = Folha(ws)
    f.larguras({"A": 4, "B": 34, "C": 40, "D": 16, "E": 22, "F": 12, "G": 14, "H": 16, "I": 16, "J": 16, "K": 16, "L": 16})
    f.titulo(1, "ALTERNATIVAS — CAPEX, CONSUMO, OPEX E ECONOMIA NO ANO 1", 12)
    f.cabecalho(3, ["#", "Alternativa", "Descrição", "CAPEX (R$)", "Origem do CAPEX", "Redução de consumo", "Geração FV (kWh/ano)",
                    "Consumo (kWh/ano)", "Manutenção (R$/ano)", "Economia energia (R$)", "Economia manutenção (R$)", "Economia ano 1 (R$)"])
    ra = {}
    for n, a in enumerate(alts):
        r = 4 + n
        f.c(r, 1, n + 1, INT, alinhar="center")
        f.c(r, 2, a["nome"], negrito=True, quebra=True)
        f.c(r, 3, a["descricao"], quebra=True, tam=9)
        f.entrada(r, 4, round(a["capex"], 2), MOEDA)
        f.c(r, 5, a["capex_origem"], tam=9, cor=CINZA)
        f.entrada(r, 6, a["reducao_pct"], PCT)
        f.entrada(r, 7, a["geracao_kwh"], KWH)
        f.c(r, 8, f"=MAX({ref['cb']}*(1-F{r})-G{r},0)", KWH)
        f.entrada(r, 9, round(a["manutencao_ano"], 2), MOEDA)
        f.c(r, 10, f"=({ref['cb']}-H{r})*{ref['tarifa']}", MOEDA)
        f.c(r, 11, f"={ref['manut']}-I{r}", MOEDA)
        f.c(r, 12, f"=J{r}+K{r}", MOEDA, negrito=True)
        ws.row_dimensions[r].height = 30
        ra[n] = r
    f.nota(4 + len(alts) + 1, "CAPEX vinculado ao preço de uma opção da proposta entra como valor (o vínculo vive no app). "
                              "Paliativos não entram no fluxo incremental: são executados em qualquer alternativa.", 12)
    ws.freeze_panes = "C4"
    _impressao(ws)

    # --------------------------------------------------------------- Fluxo de caixa
    ws = wb.create_sheet("FluxoCaixa")
    f = Folha(ws)
    f.titulo(1, "FLUXO DE CAIXA INCREMENTAL POR ALTERNATIVA (R$)", min(4 + 6 * max(len(alts), 1), 40))
    f.nota(2, "Anos além do horizonte valem zero. Descontado = fluxo × fator de desconto (TMA).", 12)
    f.cabecalho(4, ["Ano", "Ativo", "Fator de desconto", "OPEX atual (R$)"])
    f.c(3, 1, None, borda=False)
    y0 = 5
    y_fim = y0 + HORIZONTE_MAX
    for t in range(HORIZONTE_MAX + 1):
        r = y0 + t
        f.c(r, 1, t, INT, alinhar="center")
        f.c(r, 2, f"=IF(A{r}<={ref['h']},1,0)", INT, alinhar="center")
        f.c(r, 3, f"=1/(1+{ref['tma']})^A{r}", "0.0000")
        f.c(r, 4, "=0" if t == 0 else
            f"=B{r}*({ref['energia_base']}*(1+{ref['ge']})^(A{r}-1)+{ref['manut']}*(1+{ref['gm']})^(A{r}-1))", MOEDA)
    blocos = {}
    for n, a in enumerate(alts):
        c0 = 5 + 6 * n
        ws.merge_cells(start_row=3, start_column=c0, end_row=3, end_column=c0 + 5)
        f.c(3, c0, a["nome"], negrito=True, cor=AZUL, alinhar="center")
        f.cabecalho(4, ["Economia energia", "Economia manutenção", "Fluxo", "Acumulado", "Descontado", "Acumulado desc."], col=c0)
        e, m, fl, ac, ds, ad = (L(c0 + k) for k in range(6))
        ar = ra[n]
        for t in range(HORIZONTE_MAX + 1):
            r = y0 + t
            if t == 0:
                f.c(r, c0, 0, MOEDA)
                f.c(r, c0 + 1, 0, MOEDA)
                f.vinculo(r, c0 + 2, f"=-Alternativas!$D${ar}", MOEDA)
                f.c(r, c0 + 3, f"={fl}{r}", MOEDA)
                f.c(r, c0 + 4, f"={fl}{r}*$C{r}", MOEDA)
                f.c(r, c0 + 5, f"={ds}{r}", MOEDA)
            else:
                f.vinculo(r, c0, f"=$B{r}*Alternativas!$J${ar}*(1+{ref['ge']})^($A{r}-1)", MOEDA)
                f.vinculo(r, c0 + 1, f"=$B{r}*Alternativas!$K${ar}*(1+{ref['gm']})^($A{r}-1)", MOEDA)
                f.c(r, c0 + 2, f"={e}{r}+{m}{r}", MOEDA)
                f.c(r, c0 + 3, f"={ac}{r - 1}+{fl}{r}", MOEDA)
                f.c(r, c0 + 4, f"={fl}{r}*$C{r}", MOEDA)
                f.c(r, c0 + 5, f"={ad}{r - 1}+{ds}{r}", MOEDA)
        blocos[n] = {k: f"FluxoCaixa!${col}${y0}:${col}${y_fim}" for k, col in
                     (("e", e), ("m", m), ("fluxo", fl), ("acum", ac), ("desc", ds), ("acum_desc", ad))}
        blocos[n]["e1"] = f"FluxoCaixa!${e}${y0 + 1}:${e}${y_fim}"
        blocos[n]["m1"] = f"FluxoCaixa!${m}${y0 + 1}:${m}${y_fim}"
        blocos[n]["col_acum"] = c0 + 3
        for k in range(6):
            ws.column_dimensions[L(c0 + k)].width = 14
    fator1 = f"FluxoCaixa!$C${y0 + 1}:$C${y_fim}"
    opex_base = f"FluxoCaixa!$D${y0 + 1}:$D${y_fim}"
    f.larguras({"A": 6, "B": 6, "C": 10, "D": 16})
    ws.freeze_panes = "E5"
    if alts:
        graf = LineChart()
        graf.title = "Fluxo acumulado (R$)"
        graf.y_axis.title = "R$"
        graf.x_axis.title = "Ano"
        graf.height, graf.width = 9, 22
        h = int(fin["horizonte"])
        for n in range(len(alts)):
            graf.add_data(Reference(ws, min_col=blocos[n]["col_acum"], min_row=y0, max_row=y0 + h), titles_from_data=False)
        for n, serie in enumerate(graf.series):
            serie.tx = SeriesLabel(v=alts[n]["nome"])
        graf.set_categories(Reference(ws, min_col=1, min_row=y0, max_row=y0 + h))
        ws.add_chart(graf, f"A{y_fim + 3}")
    _impressao(ws)

    # --------------------------------------------------------------- Resumo (indicadores)
    f = Folha(ws_res)
    n_a = max(len(alts), 1)
    ate = 2 + n_a
    f.larguras({"A": 48, "B": 4, **{L(3 + k): 24 for k in range(n_a)}})
    _logo(ws_res, "A1")
    for linha in range(1, 5):
        ws_res.row_dimensions[linha].height = 18
    org = _org()
    f.titulo(6, (v.get("titulo") or "ESTUDO DE VIABILIDADE").upper(), ate)
    info = [("Cliente", proposta.get("cliente", {}).get("razao", "")), ("Objeto", proposta.get("objeto", "")),
            ("Responsável técnico", f"{org.get('resp', '').title()} — {org.get('registros', '')}")]
    r = 7
    for rot, val in info:
        f.c(r, 1, rot, negrito=True, borda=False)
        ws_res.merge_cells(start_row=r, start_column=3, end_row=r, end_column=ate)
        f.c(r, 3, val, borda=False, quebra=True)
        r += 1
    r += 1
    f.secao(r, "Sistema atual × projeto original", ate)
    r += 1
    for rot, chave, fmt in (("Consumo atual (kWh/ano)", "cb", KWH), ("Consumo previsto no projeto (kWh/ano)", "cp", KWH),
                            ("Consumo acima do projeto", "excesso", PCT), ("Custo anual da perda de desempenho", "perda", MOEDA),
                            ("Custo anual de energia atual", "energia_base", MOEDA)):
        f.c(r, 1, rot)
        f.vinculo(r, 3, "=" + ref[chave], fmt)
        r += 1
    f.c(r, 1, "LCC do sistema atual no horizonte (paliativos + OPEX a valor presente)", negrito=True)
    f.c(r, 3, f"={ref['paliativos']}+SUMPRODUCT({opex_base},{fator1})", MOEDA, negrito=True, fundo=FUNDO_CAB)
    lcc_base = f"$C${r}"
    r += 2
    f.cabecalho(r, ["Indicador", ""] + [a["nome"] for a in alts])
    r += 1
    pos = {}
    indicadores = [
        ("capex", "CAPEX", MOEDA), ("econ1", "Economia no ano 1", MOEDA), ("vp", "Valor presente das economias", MOEDA),
        ("vpe", "  VP da economia de energia", MOEDA), ("vpm", "  VP da economia de manutenção", MOEDA),
        ("vpl", "VPL", MOEDA), ("tir", "TIR (a.a.)", PCT), ("pb", "Payback simples", ANOS),
        ("pbd", "Breakeven econômico (payback descontado)", ANOS), ("bk_meses", "Breakeven (meses)", INT),
        ("econ_min", "Economia mínima no ano 1 para VPL = 0", MOEDA), ("margem", "Margem de segurança da economia", PCT),
        ("lcc", "LCC no horizonte (paliativos + CAPEX + OPEX a VP)", MOEDA), ("vs_proj", "Consumo × projeto original", PCT),
        ("classe", "CLASSIFICAÇÃO", None),
    ]
    for chave, *_ in indicadores:
        pos[chave] = r
        r += 1
    for k, a in enumerate(alts):
        col = L(3 + k)
        bl, ar = blocos[k], ra[k]
        P = {c: f"{col}{p}" for c, p in pos.items()}
        n_neg = f"COUNTIF({bl['acum']},\"<0\")"
        n_neg_d = f"COUNTIF({bl['acum_desc']},\"<0\")"
        formulas = {
            "capex": f"=Alternativas!$D${ar}",
            "econ1": f"=Alternativas!$L${ar}",
            "vpe": f"=SUMPRODUCT({bl['e1']},{fator1})",
            "vpm": f"=SUMPRODUCT({bl['m1']},{fator1})",
            "vp": f"={P['vpe']}+{P['vpm']}",
            "vpl": f"={P['vp']}-{P['capex']}",
            # o IRR parte de 10% e não converge em TIR muito negativa: tenta outras estimativas antes de desistir
            "tir": (f"=IF({P['econ1']}>0,IFERROR(IRR({bl['fluxo']}),IFERROR(IRR({bl['fluxo']},-0.2),"
                    f"IFERROR(IRR({bl['fluxo']},-0.5),\"—\"))),\"—\")"),
            "pb": (f"=IF({P['capex']}<=0,0,IF({P['econ1']}<=0,\"—\",IF({n_neg}>{ref['h']},\"> horizonte\","
                   f"{n_neg}-1+(-INDEX({bl['acum']},{n_neg}))/INDEX({bl['fluxo']},{n_neg}+1))))"),
            "pbd": (f"=IF({P['capex']}<=0,0,IF({P['econ1']}<=0,\"—\",IF({n_neg_d}>{ref['h']},\"> horizonte\","
                    f"{n_neg_d}-1+(-INDEX({bl['acum_desc']},{n_neg_d}))/INDEX({bl['desc']},{n_neg_d}+1))))"),
            "bk_meses": f"=IF(ISNUMBER({P['pbd']}),ROUNDUP({P['pbd']}*12,0),\"—\")",
            "econ_min": f"=IF({P['vp']}>0,{P['econ1']}*{P['capex']}/{P['vp']},\"—\")",
            "margem": f"=IF({P['capex']}>0,{P['vp']}/{P['capex']}-1,\"—\")",
            "lcc": f"={lcc_base}+{P['capex']}-{P['vp']}",
            "vs_proj": f"=IF({ref['cp']}>0,Alternativas!$H${ar}/{ref['cp']}-1,\"—\")",
            "classe": (f"=IF({P['econ1']}<=0,\"Sem economia\",IF(AND(ISNUMBER({P['pb']}),{P['pb']}<=2),\"Quick win\","
                       f"IF({P['vpl']}>0,\"Viável\",\"Inviável financeiramente\")))"),
        }
        for chave, rot, fmt in indicadores:
            forte = rot.isupper()
            if k == 0:
                f.c(pos[chave], 1, rot, negrito=forte or chave in ("vpl", "pbd"), fundo=FUNDO_CAB if forte else None)
                f.c(pos[chave], 2, None, fundo=FUNDO_CAB if forte else None)
            f.c(pos[chave], 3 + k, formulas[chave], fmt, negrito=forte or chave in ("vpl", "pbd"),
                fundo=FUNDO_CAB if forte else None, alinhar="right")
    r += 1
    f.nota(r, "Payback simples: anos até a economia acumulada igualar o CAPEX. Breakeven econômico: anos até o VPL "
              "acumulado zerar (descontado pela TMA). Economia mínima: economia do ano 1 que zera o VPL no horizonte. "
              "Valores ilustrativos até a validação com medições e cotações.", ate)
    f.nota(r + 1, "Fundo amarelo = entrada editável; verde = vínculo com outra aba.", ate)
    ws_res.freeze_panes = "A7"
    _impressao(ws_res, paisagem=n_a > 2)
    ws_res.page_setup.fitToHeight = 1

    # --------------------------------------------------------------- Sensibilidade
    ws = wb.create_sheet("Sensibilidade")
    f = Folha(ws)
    f.larguras({"A": 34, **{L(2 + k): 15 for k in range(len(VARIACOES_ECONOMIA))}})
    f.titulo(1, "SENSIBILIDADE DO VPL (R$)", 1 + len(VARIACOES_ECONOMIA))
    r = 3
    for titulo, variacoes, termo in (
            ("Variação da economia anual", VARIACOES_ECONOMIA, lambda vp, vpe, vpm, cap, d: f"={vp}*(1+{d})-{cap}"),
            ("Variação do CAPEX", VARIACOES_CAPEX, lambda vp, vpe, vpm, cap, d: f"={vp}-{cap}*(1+{d})"),
            ("Variação da tarifa de energia", VARIACOES_TARIFA, lambda vp, vpe, vpm, cap, d: f"={vpe}*(1+{d})+{vpm}-{cap}")):
        f.secao(r, titulo, 1 + len(variacoes))
        r += 1
        hr = r      # linha com as variações
        f.c(hr, 1, "Alternativa", negrito=True, cor=AZUL, fundo=FUNDO_CAB)
        for j, d in enumerate(variacoes):
            f.c(hr, 2 + j, d, PCT, negrito=True, cor=AZUL, fundo=FUNDO_CAB, alinhar="center")
        r += 1
        for k, a in enumerate(alts):
            col = L(3 + k)
            vp, vpe, vpm, cap = (f"Resumo!${col}${pos[c]}" for c in ("vp", "vpe", "vpm", "capex"))
            f.c(r, 1, a["nome"])
            for j in range(len(variacoes)):
                f.vinculo(r, 2 + j, termo(vp, vpe, vpm, cap, f"{L(2 + j)}${hr}"), MOEDA)
            r += 1
        r += 1
    _impressao(ws)

    wb.properties.title = v.get("titulo") or "Estudo de viabilidade"
    wb.properties.creator = org.get("fantasia", "KAPTEK")
    wb.calculation.fullCalcOnLoad = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
