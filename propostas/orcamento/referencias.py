"""
Bases de referência: preços (EMOP-RJ, SCO-Rio, SINAPI, SICRO, ORSE, cotações...), pisos de
convenções coletivas (SENGE, SINDRATAR, SINDISTAL/SINTRAINDISTAL, SINDUSCON/SINTRACONST...),
CUB, encargos sociais e índices de reajuste.

Os arquivos oficiais são baixados pelo usuário e importados (XLSX, CSV ou PDF). Nada vem
embutido: cada valor fica gravado com fonte, código, data-base e UF para rastreabilidade.

Uso no orçamento:
- busca em todas as bases, na ordem da hierarquia de fontes (RJ: EMOP → SCO-Rio → SINAPI → SICRO,
  Decreto Estadual 48.929/2024), para compor fornecimentos e subcontratações;
- cesta de preços por item (mediana, média ou menor valor, com alerta de valores fora da faixa);
- atualização da data-base por índice (INCC, IPCA...) quando a série estiver importada;
- custo/h dos perfis a partir do piso da convenção + adicionais + encargos + benefícios.
"""

import csv
import io
import json
import os
import re
import sqlite3
import statistics
import unicodedata
import datetime as dt

from . import parametros as pm

RAIZ = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
CAMINHO_DB = os.path.join(RAIZ, "dados", "referencias.sqlite")
CAMINHO_FONTES = os.path.join(RAIZ, "dados", "fontes.json")

ESQUEMA = """
CREATE TABLE IF NOT EXISTS itens (
    id INTEGER PRIMARY KEY, fonte TEXT, codigo TEXT, descricao TEXT, busca TEXT, unidade TEXT, valor REAL,
    data_base TEXT, uf TEXT, desonerado INTEGER, tipo TEXT, arquivo TEXT, importado_em TEXT);
CREATE INDEX IF NOT EXISTS ix_itens_fonte ON itens(fonte, uf, data_base);
CREATE INDEX IF NOT EXISTS ix_itens_codigo ON itens(codigo);
CREATE TABLE IF NOT EXISTS pisos (
    id INTEGER PRIMARY KEY, entidade TEXT, cct TEXT, funcao TEXT, piso REAL, carga_semanal REAL,
    data_base TEXT, uf TEXT, obs TEXT, importado_em TEXT);
CREATE TABLE IF NOT EXISTS cub (
    id INTEGER PRIMARY KEY, fonte TEXT, uf TEXT, padrao TEXT, mes TEXT, valor REAL, importado_em TEXT);
CREATE TABLE IF NOT EXISTS indices (
    id INTEGER PRIMARY KEY, nome TEXT, mes TEXT, valor REAL, UNIQUE(nome, mes));
CREATE TABLE IF NOT EXISTS encargos (
    id INTEGER PRIMARY KEY, fonte TEXT, uf TEXT, regime TEXT, desonerado INTEGER, pct REAL, data_base TEXT);
"""

# palavras que identificam cada coluna no cabeçalho (comparadas sem acento e em minúsculas)
PALAVRAS = {
    "codigo": ("codigo", "cod.", "cod ", "item", "referencia"),
    "descricao": ("descricao", "discriminacao", "especificacao", "denominacao", "servico", "insumo", "funcao", "cargo"),
    "unidade": ("unidade", "unid", "und", "un."),
    "valor": ("preco", "custo", "valor", "mediana", "piso", "salario", "total"),
}

# séries de índices: mês na primeira coluna e número-índice ou variação na outra
PALAVRAS_INDICES = {
    "codigo": ("mes", "data", "periodo", "competencia", "referencia"),
    "valor": ("indice", "variacao", "numero", "valor"),
}

# presets de importação por fonte; "valor_preferido" desempata quando há mais de uma coluna de preço
PRESETS = {
    "EMOP-RJ": {"tipo": "composicao", "valor_preferido": ("preco unitario", "custo unitario", "preco", "custo")},
    "SCO-RIO": {"tipo": "composicao", "valor_preferido": ("preco unitario", "custo unitario", "preco", "custo")},
    "SINAPI": {"tipo": "insumo", "valor_preferido": ("preco mediano", "custo total", "preco", "custo")},
    "SICRO": {"tipo": "composicao", "valor_preferido": ("custo unitario", "preco unitario", "custo", "preco")},
    "ABRAVA": {"tipo": "referencia", "valor_preferido": ("valor", "preco", "custo")},
    "Genérico": {"tipo": "referencia", "valor_preferido": ("preco", "custo", "valor", "total")},
}

# Lei 4.950-A/1966: 6 salários mínimos para 6 h/dia; cada hora além de 6 vale 25% a mais
SALARIO_MINIMO_2026 = 1621.00        # Decreto 12.797/2025
BASE_CONGELADA_STF = 1212.00         # ADPF 53, 149 e 171: base fixada no salário mínimo de 2022


# ----------------------------------------------------------------------------- utilitários
def normalizar(texto):
    t = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"\s+", " ", t).strip()


def converter_valor(x):
    """'R$ 1.234,56', '1234.56', 1234.56 → float; vazio ou texto → None."""
    if x is None:
        return None
    if isinstance(x, (int, float)):
        return float(x)
    t = re.sub(r"[^\d,.\-]", "", str(x))
    if not t or not re.search(r"\d", t):
        return None
    if "," in t and "." in t:
        t = t.replace(".", "").replace(",", ".") if t.rfind(",") > t.rfind(".") else t.replace(",", "")
    elif "," in t:
        t = t.replace(".", "").replace(",", ".")
    try:
        return float(t)
    except ValueError:
        return None


def info_do_nome(nome):
    """UF e data-base (AAAA-MM) a partir do nome do arquivo, como SINAPI_..._RJ_202405_Desonerado.xlsx."""
    base = os.path.basename(nome or "")
    uf = next((m for m in re.findall(r"(?<![A-Z])([A-Z]{2})(?![A-Z])", base.upper())
               if m in {"AC", "AL", "AM", "AP", "BA", "CE", "DF", "ES", "GO", "MA", "MG", "MS", "MT", "PA", "PB",
                        "PE", "PI", "PR", "RJ", "RN", "RO", "RR", "RS", "SC", "SE", "SP", "TO"}), "")
    mes = ""
    m = re.search(r"(20\d{2})[-_.]?(0[1-9]|1[0-2])(?!\d)", base)
    if m:
        mes = f"{m.group(1)}-{m.group(2)}"
    else:
        m = re.search(r"(0[1-9]|1[0-2])[-_.](20\d{2})", base)
        if m:
            mes = f"{m.group(2)}-{m.group(1)}"
    desonerado = None
    if re.search(r"n[aã]o[ _-]?desonerad", base, re.I):
        desonerado = False
    elif re.search(r"desonerad", base, re.I):
        desonerado = True
    return {"uf": uf, "data_base": mes, "desonerado": desonerado}


MESES = {"jan": 1, "fev": 2, "mar": 3, "abr": 4, "mai": 5, "jun": 6, "jul": 7, "ago": 8, "set": 9, "out": 10, "nov": 11, "dez": 12}


def converter_mes(x):
    """'2024-05', '05/2024', 'mai/2024', 'mai/24' ou data → 'AAAA-MM'; não reconhecido → ''."""
    if isinstance(x, (dt.date, dt.datetime)):
        return f"{x.year:04d}-{x.month:02d}"
    t = normalizar(x)
    m = re.match(r"^(\d{4})[-/.](\d{1,2})", t)
    if m:
        return f"{int(m.group(1)):04d}-{int(m.group(2)):02d}"
    m = re.match(r"^(\d{1,2})[-/.](\d{4})", t)
    if m:
        return f"{int(m.group(2)):04d}-{int(m.group(1)):02d}"
    m = re.match(r"^([a-z]{3})[a-z]*[-/. ]?(\d{2,4})$", t)
    if m and m.group(1) in MESES:
        ano = int(m.group(2))
        ano = ano + 2000 if ano < 100 else ano
        return f"{ano:04d}-{MESES[m.group(1)]:02d}"
    return ""


def mes_anterior(mes, n):
    a, m = (int(x) for x in mes.split("-"))
    total = a * 12 + (m - 1) - n
    return f"{total // 12:04d}-{total % 12 + 1:02d}"


def meses_entre(inicio, fim):
    a1, m1 = (int(x) for x in inicio.split("-"))
    a2, m2 = (int(x) for x in fim.split("-"))
    return (a2 - a1) * 12 + (m2 - m1)


# ----------------------------------------------------------------------------- leitura de arquivos
def ler_tabela(conteudo, nome, aba=None):
    """Linhas (listas de células) de um XLSX, CSV ou PDF. No PDF, junta as tabelas de todas as páginas."""
    ext = os.path.splitext(nome.lower())[1]
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
        ws = wb[aba] if aba else wb.worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    if ext in (".csv", ".txt"):
        texto = conteudo.decode("utf-8-sig", errors="replace")
        try:
            dialeto = csv.Sniffer().sniff(texto[:5000], delimiters=";,\t|")
        except csv.Error:
            dialeto = csv.excel
            dialeto.delimiter = ";" if texto.count(";") > texto.count(",") else ","
        return [r for r in csv.reader(io.StringIO(texto), dialeto)]
    if ext == ".pdf":
        import pdfplumber
        linhas = []
        with pdfplumber.open(io.BytesIO(conteudo)) as pdf:
            for pagina in pdf.pages:
                for tabela in pagina.extract_tables() or []:
                    linhas += [[(c or "").replace("\n", " ").strip() for c in row] for row in tabela]
                pagina.close()          # sem isso cada página fica na memória (~5 MB): um PDF de 500 páginas estoura
        return linhas
    raise ValueError(f"formato não suportado: {ext} (use XLSX, CSV ou PDF)")


def abas_xlsx(conteudo):
    from openpyxl import load_workbook
    return load_workbook(io.BytesIO(conteudo), read_only=True).sheetnames


def detectar_cabecalho(linhas, preset="Genérico", max_linhas=60, uf=None, palavras=None):
    """Acha a linha de cabeçalho (a que mais casa com as palavras-chave) e o mapa campo → coluna.

    Com `uf`, uma coluna cujo cabeçalho é a própria UF ("RJ") vira o valor: é o layout do SINAPI a
    partir de 2025, com o preço de todas as UFs no mesmo arquivo, uma coluna por UF."""
    pref = PRESETS.get(preset, PRESETS["Genérico"])["valor_preferido"]
    palavras = palavras or PALAVRAS
    alvo_uf = normalizar(uf) if uf else ""
    melhor = (-1, None, {})
    for i, row in enumerate(linhas[:max_linhas]):
        celulas = [normalizar(c) for c in row]
        mapa = {}
        for campo, chaves in palavras.items():
            candidatos = [j for j, c in enumerate(celulas) if c and any(p in c + " " for p in chaves)]
            if not candidatos:
                continue
            if campo == "valor":
                ordenados = sorted(candidatos, key=lambda j: next((k for k, p in enumerate(pref) if p in celulas[j]), len(pref)))
                mapa[campo] = ordenados[0]
            else:
                livres = [j for j in candidatos if j not in mapa.values()]
                if livres:
                    mapa[campo] = livres[0]
        if alvo_uf and alvo_uf in celulas:
            mapa["valor"] = celulas.index(alvo_uf)
        pontos = len(mapa)
        if pontos > melhor[0]:
            melhor = (pontos, i, mapa)
    if melhor[0] < 2:
        return None, {}
    return melhor[1], melhor[2]


def extrair_itens(linhas, inicio, mapa):
    """Itens a partir da linha seguinte ao cabeçalho; ignora linhas sem descrição ou sem valor numérico."""
    itens = []
    for row in linhas[inicio + 1:]:
        def cel(campo):
            j = mapa.get(campo)
            return row[j] if j is not None and j < len(row) else None
        desc = str(cel("descricao") or "").strip()
        valor = converter_valor(cel("valor"))
        if not desc or valor is None:
            continue
        itens.append({"codigo": str(cel("codigo") or "").strip(), "descricao": desc,
                      "unidade": str(cel("unidade") or "").strip(), "valor": valor})
    return itens


# ----------------------------------------------------------------------------- banco
def conectar(caminho=None):
    """Abre (e cria, se preciso) o banco. KAPTEK_REFERENCIAS aponta para outro arquivo (testes, base compartilhada)."""
    # o Streamlit pode reexecutar o script em outra thread; o acesso é sempre sequencial
    con = sqlite3.connect(caminho or os.environ.get("KAPTEK_REFERENCIAS") or CAMINHO_DB, check_same_thread=False)
    con.row_factory = sqlite3.Row
    con.executescript(ESQUEMA)
    return con


def _agora():
    return pm.agora()


def importar_itens(con, itens, fonte, data_base="", uf="", desonerado=None, tipo="", arquivo="", substituir=True):
    """Grava itens de preço. Com `substituir`, apaga antes a mesma fonte + UF + data-base."""
    if substituir:
        con.execute("DELETE FROM itens WHERE fonte=? AND uf=? AND data_base=?", (fonte, uf, data_base))
    con.executemany(
        "INSERT INTO itens (fonte, codigo, descricao, busca, unidade, valor, data_base, uf, desonerado, tipo, arquivo, importado_em)"
        " VALUES (?,?,?,?,?,?,?,?,?,?,?,?)",
        [(fonte, i["codigo"], i["descricao"], normalizar(f"{i['codigo']} {i['descricao']}"), i["unidade"], i["valor"],
          data_base, uf, None if desonerado is None else int(desonerado), tipo, arquivo, _agora()) for i in itens])
    con.commit()
    return len(itens)


def fontes_importadas(con):
    return [dict(r) for r in con.execute(
        "SELECT fonte, uf, data_base, COUNT(*) AS itens, MAX(importado_em) AS importado_em FROM itens "
        "GROUP BY fonte, uf, data_base ORDER BY fonte, uf, data_base DESC")]


def apagar_fonte(con, fonte, uf, data_base):
    con.execute("DELETE FROM itens WHERE fonte=? AND uf=? AND data_base=?", (fonte, uf, data_base))
    con.commit()


def ordem_hierarquia(hierarquia, uf):
    """Lista de fontes em ordem de prioridade para a UF (padrão quando a UF não tem regra própria)."""
    return list((hierarquia or {}).get(uf) or (hierarquia or {}).get("padrao") or [])


def buscar(con, termo, uf=None, fontes=None, hierarquia=None, limite=60):
    """Itens que contêm todas as palavras do termo (código ou descrição), na ordem da hierarquia."""
    palavras = [p for p in normalizar(termo).split() if p]
    if not palavras:
        return []
    sql = "SELECT * FROM itens WHERE " + " AND ".join("busca LIKE ?" for _ in palavras)
    args = [f"%{p}%" for p in palavras]
    if uf:
        sql += " AND (uf=? OR uf='')"
        args.append(uf)
    if fontes:
        sql += " AND fonte IN (" + ",".join("?" for _ in fontes) + ")"
        args += list(fontes)
    sql += " LIMIT ?"
    args.append(limite * 4)
    linhas = [dict(r) for r in con.execute(sql, args)]
    ordem = ordem_hierarquia(hierarquia, uf or "")
    rank = {f: k for k, f in enumerate(ordem)}
    # ordenações estáveis: descrição mais curta < data-base mais recente < prioridade da fonte
    linhas.sort(key=lambda r: len(r["descricao"]))
    linhas.sort(key=lambda r: r["data_base"] or "", reverse=True)
    linhas.sort(key=lambda r: rank.get(r["fonte"], len(rank)))
    for r in linhas:
        r["prioridade"] = rank.get(r["fonte"], None)
    return linhas[:limite]


# ----------------------------------------------------------------------------- índices e atualização
def importar_indices(con, nome, serie):
    """`serie`: lista de (AAAA-MM, valor do número-índice)."""
    con.executemany("INSERT OR REPLACE INTO indices (nome, mes, valor) VALUES (?,?,?)",
                    [(nome, m, float(v)) for m, v in serie])
    con.commit()
    return len(serie)


def indice_de_variacoes(serie):
    """Variações mensais em % [(AAAA-MM, 0,45), ...] → números-índice com base 100 no mês anterior ao
    primeiro. O Banco Central (SGS) e os boletins da FGV publicam o INCC-M e o IGP-M assim."""
    ordenada = sorted((m, float(v)) for m, v in serie)
    if not ordenada:
        return []
    repetidos = sorted({m for (a, _), (m, _) in zip(ordenada, ordenada[1:]) if a == m})
    if repetidos:
        raise ValueError("meses repetidos na série: " + ", ".join(repetidos))
    faltando = [mes_anterior(m, k) for (a, _), (m, _) in zip(ordenada, ordenada[1:])
                for k in range(meses_entre(a, m) - 1, 0, -1)]
    if faltando:
        raise ValueError("meses faltando na série: " + ", ".join(faltando))
    indice = [(mes_anterior(ordenada[0][0], 1), 100.0)]
    for m, v in ordenada:
        indice.append((m, indice[-1][1] * (1 + v / 100)))
    return indice


def indices_disponiveis(con):
    return [dict(r) for r in con.execute(
        "SELECT nome, MIN(mes) AS de, MAX(mes) AS ate, COUNT(*) AS meses FROM indices GROUP BY nome ORDER BY nome")]


def atualizar(con, valor, data_base, data_alvo, indice):
    """Valor levado da data-base para a data-alvo pelo número-índice. Devolve (valor, fator ou None)."""
    if not (con and data_base and data_alvo and indice) or data_base == data_alvo:
        return valor, (1.0 if data_base == data_alvo and data_base else None)
    q = "SELECT valor FROM indices WHERE nome=? AND mes=?"
    a = con.execute(q, (indice, data_base)).fetchone()
    b = con.execute(q, (indice, data_alvo)).fetchone()
    if not a or not b or not a["valor"]:
        return valor, None
    fator = b["valor"] / a["valor"]
    return valor * fator, fator


# ----------------------------------------------------------------------------- cesta de preços
def cesta(referencias, criterio="mediana", faixa_alerta=0.25, minimo_fontes=3, con=None, data_alvo="", indice=""):
    """Preço de referência de um item a partir de várias fontes e cotações.

    Cada referência: {fonte, codigo, valor, data_base}. Os valores são atualizados para a data-alvo
    quando há índice; valores fora de ±faixa da mediana ficam marcados. O critério escolhe o valor
    adotado entre os não marcados (mediana, média ou menor)."""
    linhas = []
    for r in referencias:
        v = converter_valor(r.get("valor"))
        if v is None:
            continue
        v_at, fator = atualizar(con, v, r.get("data_base", ""), data_alvo, indice) if con else (v, None)
        linhas.append({**r, "valor": v, "valor_atualizado": v_at, "fator": fator})
    if not linhas:
        return {"valor": None, "linhas": [], "avisos": ["Nenhuma referência com valor."]}
    med = statistics.median(l["valor_atualizado"] for l in linhas)
    for l in linhas:
        l["desvio"] = l["valor_atualizado"] / med - 1 if med else 0.0
        l["fora_da_faixa"] = abs(l["desvio"]) > faixa_alerta
    validos = [l["valor_atualizado"] for l in linhas if not l["fora_da_faixa"]] or [l["valor_atualizado"] for l in linhas]
    adotado = {"mediana": statistics.median, "media": statistics.mean, "menor": min}[criterio](validos)
    avisos = []
    if len(linhas) < minimo_fontes:
        avisos.append(f"Só {len(linhas)} referência(s); recomenda-se no mínimo {minimo_fontes}.")
    fora = [l["fonte"] for l in linhas if l["fora_da_faixa"]]
    if fora:
        avisos.append(f"Fora de ±{faixa_alerta:.0%} da mediana (desconsideradas): {', '.join(fora)}.")
    sem_atualizar = [l["fonte"] for l in linhas if indice and l.get("data_base") and l["data_base"] != data_alvo and l["fator"] is None]
    if sem_atualizar:
        avisos.append(f"Sem índice {indice} para atualizar: {', '.join(sem_atualizar)}.")
    return {"valor": adotado, "criterio": criterio, "mediana": med,
            "media": statistics.mean(l["valor_atualizado"] for l in linhas),
            "menor": min(l["valor_atualizado"] for l in linhas), "n": len(linhas), "n_validos": len(validos),
            "linhas": linhas, "avisos": avisos}


# ----------------------------------------------------------------------------- convenções, encargos e CUB
def importar_pisos(con, pisos, entidade, cct="", data_base="", uf=""):
    """`pisos`: lista de {funcao, piso, carga_semanal?, obs?}."""
    con.executemany(
        "INSERT INTO pisos (entidade, cct, funcao, piso, carga_semanal, data_base, uf, obs, importado_em) VALUES (?,?,?,?,?,?,?,?,?)",
        [(entidade, cct, p["funcao"], float(p["piso"]), float(p.get("carga_semanal") or 44), data_base, uf,
          p.get("obs", ""), _agora()) for p in pisos])
    con.commit()
    return len(pisos)


def listar_pisos(con, entidade=None):
    sql, args = "SELECT * FROM pisos", []
    if entidade:
        sql += " WHERE entidade=?"
        args.append(entidade)
    return [dict(r) for r in con.execute(sql + " ORDER BY entidade, funcao", args)]


def piso_lei_4950a(salario_base, horas_dia):
    """Salário mínimo profissional de engenheiro (Lei 4.950-A/1966): 6 SM para 6 h/dia, +25% por hora além de 6."""
    horas = max(6.0, float(horas_dia))
    return float(salario_base) * (6 + 1.25 * (horas - 6))


def custo_hora_piso(piso, adicional_pct=0.0, encargos_pct=0.0, beneficios=0.0, horas_produtivas=140.0):
    """(piso × (1 + adicionais) × (1 + encargos) + benefícios) ÷ horas produtivas do mês."""
    if horas_produtivas <= 0:
        return 0.0
    return (float(piso) * (1 + float(adicional_pct)) * (1 + float(encargos_pct)) + float(beneficios)) / float(horas_produtivas)


def importar_encargos(con, linhas):
    """`linhas`: lista de {fonte, uf, regime (horista|mensalista), desonerado, pct, data_base}."""
    con.executemany("INSERT INTO encargos (fonte, uf, regime, desonerado, pct, data_base) VALUES (?,?,?,?,?,?)",
                    [(l.get("fonte", "SINAPI"), l["uf"], l["regime"], int(bool(l.get("desonerado"))), float(l["pct"]),
                      l.get("data_base", "")) for l in linhas])
    con.commit()
    return len(linhas)


def encargos(con, uf, regime="mensalista", desonerado=False):
    r = con.execute("SELECT pct FROM encargos WHERE uf=? AND regime=? AND desonerado=? ORDER BY data_base DESC LIMIT 1",
                    (uf, regime, int(bool(desonerado)))).fetchone()
    return r["pct"] if r else None


def importar_cub(con, linhas):
    """`linhas`: lista de {fonte, uf, padrao, mes, valor}."""
    con.executemany("INSERT INTO cub (fonte, uf, padrao, mes, valor, importado_em) VALUES (?,?,?,?,?,?)",
                    [(l.get("fonte", "SINDUSCON"), l["uf"], l["padrao"], l["mes"], float(l["valor"]), _agora()) for l in linhas])
    con.commit()
    return len(linhas)


def listar_cub(con, uf=None):
    sql, args = "SELECT * FROM cub", []
    if uf:
        sql += " WHERE uf=?"
        args.append(uf)
    return [dict(r) for r in con.execute(sql + " ORDER BY mes DESC, padrao", args)]


# ----------------------------------------------------------------------------- registro de entidades
def carregar_fontes(caminho=None):
    with open(caminho or CAMINHO_FONTES, encoding="utf-8") as f:
        return json.load(f)
