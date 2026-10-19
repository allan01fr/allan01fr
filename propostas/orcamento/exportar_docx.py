"""
Proposta comercial no template aprovado da KAPTEK (modelos/proposta.docx).

Regra da marca: preencher o template, nunca remontar o documento. Este módulo só
substitui os campos entre colchetes e ajusta o número de linhas das tabelas e
dos itens de lista (escopo, exclusões, premissas, prazos, honorários,
pagamento), clonando a linha ou o parágrafo do próprio template.

Com várias fases, os honorários, os prazos e o pagamento saem por fase, e as fases de obra
trazem as regras do % sobre as medições, do pacote de horas e da taxa de desempenho. A
estimativa interna do valor global não gera proposta.

`substituir()` vem de kaptek-doc/scripts/preencher.py (habilidade KAPTEK).
"""

import copy
import io
import os
import re
import unicodedata

from docx import Document
from docx.text.paragraph import Paragraph

from . import catalogo as ct

RAIZ = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
TEMPLATE = os.path.join(RAIZ, "modelos", "proposta.docx")

FRASE_DESLOCAMENTO = ("Não estão inclusos deslocamentos fora da Região Metropolitana do Rio de Janeiro, "
                      "que serão orçados à parte mediante aprovação prévia.")
PREMISSAS_FINAIS = [
    "Os prazos são contados em dias corridos a partir do cumprimento cumulativo das premissas acima.",
    "Atraso na liberação de acesso ou de documentação suspende a contagem dos prazos, sem ônus para a contratada.",
]


# ----------------------------------------------------------------------------- texto
def substituir(par, de, para):
    """Troca texto no parágrafo preservando a formatação (de kaptek-doc/scripts/preencher.py).
    Se o trecho estiver fragmentado entre runs, consolida no primeiro run."""
    if de not in par.text or not par.runs:
        return False
    for r in par.runs:
        if de in r.text:
            r.text = r.text.replace(de, para)
            return True
    novo = par.text.replace(de, para)
    par.runs[0].text = novo
    for r in par.runs[1:]:
        r.text = ""
    return True


def _definir(par, texto):
    """Troca todo o texto de um parágrafo mantendo a formatação do primeiro run."""
    if not par.runs:
        par.add_run(texto)
        return
    par.runs[0].text = texto
    for r in par.runs[1:]:
        r.text = ""


def _definir_item(par, texto):
    """Item de lista do template: run 0 é o travessão cobre, run 1 o texto grafite."""
    if len(par.runs) >= 2:
        par.runs[1].text = texto
        for r in par.runs[2:]:
            r.text = ""
    else:
        _definir(par, "—   " + texto)


def _todos_paragrafos(doc):
    yield from doc.paragraphs
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                yield from c.paragraphs


def moeda(v):
    s = f"{v:,.2f}"
    return "R$ " + s.replace(",", "§").replace(".", ",").replace("§", ".")


def pct_br(p):
    v = round(p * 100, 1)
    return (f"{v:.0f}" if v == int(v) else f"{v:.1f}".replace(".", ",")) + "%"


def _data_br(iso):
    try:
        a, m, d = iso.split("-")
        return f"{d}/{m}/{a}"
    except (ValueError, AttributeError):
        return str(iso)


def slug(texto):
    t = unicodedata.normalize("NFKD", texto or "").encode("ascii", "ignore").decode()
    t = re.sub(r"[^A-Za-z0-9]+", "-", t).strip("-")
    return t[:40] or "cliente"


# ----------------------------------------------------------------------------- estrutura
def _listas_por_secao(doc):
    """Itens de lista (parágrafos que começam com travessão) agrupados pelo título da seção."""
    secoes, atual = {}, None
    for p in doc.paragraphs:
        if p.style is not None and p.style.name.startswith("Heading"):
            atual = p.text.upper()
            continue
        if atual and p.text.startswith("—"):
            secoes.setdefault(atual, []).append(p)
    return secoes


def _lista(secoes, chave):
    for titulo, pars in secoes.items():
        if chave in titulo:
            return pars
    raise KeyError(f"seção '{chave}' não encontrada no template")


def _ajustar_lista(pars, textos):
    textos = list(textos) or ["Não se aplica."]
    ultimo = pars[-1]
    for i, t in enumerate(textos):
        if i < len(pars):
            alvo = pars[i]
        else:
            novo = copy.deepcopy(pars[0]._p)
            ultimo._p.addnext(novo)
            alvo = Paragraph(novo, pars[0]._parent)
            ultimo = alvo
        _definir_item(alvo, t)
    for p in pars[len(textos):]:
        p._p.getparent().remove(p._p)


def _tabela(doc, primeiro_cabecalho, segundo=None):
    for t in doc.tables:
        cab = [c.text.strip() for c in t.rows[0].cells]
        if cab and cab[0] == primeiro_cabecalho and (segundo is None or (len(cab) > 1 and cab[1] == segundo)):
            return t
    raise KeyError(f"tabela '{primeiro_cabecalho}' não encontrada no template")


def _preencher_linhas(tabela, linhas, manter_ultima=False):
    """Recria as linhas de dados clonando a primeira linha de dados do template."""
    trs = tabela._tbl.tr_lst
    proto = trs[1]
    ancora = trs[-1] if manter_ultima else None
    for tr in trs[1:len(trs) - (1 if manter_ultima else 0)]:
        if tr is not proto:
            tabela._tbl.remove(tr)
    for valores in linhas:
        novo = copy.deepcopy(proto)
        if ancora is not None:
            ancora.addprevious(novo)
        else:
            tabela._tbl.append(novo)
    tabela._tbl.remove(proto)
    inicio = 1
    for i, valores in enumerate(linhas):
        row = tabela.rows[inicio + i]
        for c, v in zip(row.cells, valores):
            if c.paragraphs:
                _definir(c.paragraphs[0], str(v))
                for extra in c.paragraphs[1:]:
                    extra._p.getparent().remove(extra._p)


# ----------------------------------------------------------------------------- conteúdo
def _rateio_servicos(opcao, res):
    """Preço dos serviços repartido por serviço, na proporção do custo direto de cada um.
    Devolve (índice do item, nome, valor); a soma fecha exatamente com o preço dos serviços."""
    pesos = []
    for i, it in enumerate(opcao["itens"]):
        mo = sum(l["custo_mo"] for l in res["linhas"] if l["item"] == i)
        dr = sum(d["valor"] for d in res["diretos"] if d["item"] == i)
        nome = next((l["servico_nome"] for l in res["linhas"] if l["item"] == i), None)
        if nome and mo + dr > 0:
            pesos.append((i, nome, mo + dr))
    total = sum(p for *_, p in pesos) or 1
    valores = [round(res["preco_servicos"] * p / total, 2) for *_, p in pesos]
    if valores:
        valores[-1] = round(res["preco_servicos"] - sum(valores[:-1]), 2)
    return [(i, n, v) for (i, n, _), v in zip(pesos, valores)]


def _h(v):
    return f"{v:,.0f}".replace(",", ".") + " h"


def _linhas_percentual(fr, prefixo):
    hp = fr["percentual"]
    pct, base = pct_br(hp["pct"]), moeda(hp["base"])
    if hp["forma"] == "fixo_global":
        return [(f"{prefixo}{pct} sobre o valor global estimado da obra ({base}), faturado conforme as medições",
                 pct, "—", moeda(hp["fixo"]))]
    linhas = []
    if hp["forma"] == "fixo_mais_medicoes" and hp["parcela_fixa"] > 0:
        linhas.append((f"{prefixo}parcela fixa (mobilização e planejamento inicial)", "1", moeda(hp["parcela_fixa"]),
                       moeda(hp["parcela_fixa"])))
    linhas.append((f"{prefixo}{pct} sobre o valor de cada medição da obra (valor global estimado em {base})",
                   pct, "—", f"≈ {moeda(hp['estimado'])} (estimado)"))
    return linhas


def _linhas_pacote(fr, prefixo):
    pk = fr["pacote"]
    if pk["forma"] == "mensal":
        desc = f"{prefixo}pacote de {pk['horas_mes']:.0f} h/mês × {pk['meses']} meses ({_h(pk['horas'])})"
    else:
        desc = f"{prefixo}pacote de {_h(pk['horas'])}, a consumir em até {pk['meses']} meses"
    total = fr["preco"] - fr["preco_fornec"]
    return [(desc, _h(pk["horas"]), moeda(total / pk["horas"]) if pk["horas"] else "—", moeda(total)),
            ("Hora excedente ao pacote (sob demanda, mediante autorização prévia)", "1 h", moeda(pk["taxa_excedente"]),
             "sob demanda")]


def _linhas_fase(cat, params, opcao, fr, prefixo="", detalhar=True):
    """Linhas de honorários de uma fase. `detalhar` (fase única) lista serviços, perfis e fornecimentos."""
    linhas = []
    cob = fr["cobranca"]
    if cob == "hora":
        if detalhar:
            for p, nome in cat["perfis"].items():
                h = fr["horas_perfil"][p]
                if h > 0:
                    taxa = float(params["perfis"][p]["taxa_venda_h"])
                    linhas.append((f"Hora técnica — {nome}", f"{h:.1f} h".replace(".", ","), moeda(taxa), moeda(h * taxa)))
            if fr["reembolso"] > 0:
                linhas.append(("Despesas reembolsáveis (estimativa, inclui ART)", "1", moeda(fr["reembolso"]), moeda(fr["reembolso"])))
        else:
            valor = fr["honorarios"]
            linhas.append((f"{prefixo}hora técnica (estimado, teto {moeda(fr['preco_hora_teto'] - fr['preco_fornec'])})",
                           _h(fr["total_h"]), moeda(valor / fr["total_h"]) if fr["total_h"] else "—", moeda(valor)))
    elif cob == "percentual":
        linhas += _linhas_percentual(fr, prefixo or "Honorários de engenharia — ")
    elif cob == "pacote":
        linhas += _linhas_pacote(fr, prefixo or "Fiscalização e planejamento — ")
    elif detalhar:
        for _, nome, v in _rateio_servicos(opcao, fr):
            linhas.append((f"Honorários de engenharia — {nome}", "1", moeda(v), moeda(v)))
    else:
        linhas.append((f"{prefixo}escopo fechado", "1 vb", moeda(fr["preco_servicos"]), moeda(fr["preco_servicos"])))
    if fr["fornecimentos"]:
        if detalhar:
            for f in fr["fornecimentos"]:
                qtd = float(f.get("qtd") or 0)
                unit = f["preco"] / qtd if qtd else f["preco"]
                q_txt = f"{qtd:g} {f.get('unid', '')}".strip().replace(".", ",")
                linhas.append((f.get("descricao", "Fornecimento"), q_txt, moeda(unit), moeda(f["preco"])))
        else:
            linhas.append((f"{prefixo}fornecimentos", "1 vb", moeda(fr["preco_fornec"]), moeda(fr["preco_fornec"])))
    if fr["desempenho"]:
        linhas.append((f"{prefixo}taxa de desempenho (variável, adicional; critérios nas condições de pagamento)",
                       "—", f"até {moeda(fr['desempenho']['teto'])}", "variável"))
    return linhas


def _linhas_honorarios(cat, params, opcao, res):
    fases = res["fases"]
    linhas = []
    if len(fases) == 1:
        linhas = _linhas_fase(cat, params, opcao, fases[0])
    else:
        for n, fr in enumerate(fases, 1):
            linhas += _linhas_fase(cat, params, opcao, fr, prefixo=f"Fase {n} — {fr['nome']}: ", detalhar=False)
    if any(d["descricao"].startswith("Taxa de ART") for d in res["diretos"]) and not \
            (len(fases) == 1 and fases[0]["cobranca"] == "hora"):
        linhas.append(("Taxa de ART junto ao CREA" + (" (uma por fase)" if len(fases) > 1 else ""), "1", "inclusa", "inclusa"))
    return linhas


def _total_texto(res):
    if len(res["fases"]) == 1 and res["fases"][0]["cobranca"] == "hora":
        return f"{moeda(res['preco'])} (estimado; teto {moeda(res['fases'][0]['preco_hora_teto'])})"
    if res["preco_estimado"] > 0:
        return f"{moeda(res['preco_fixo'])} + {moeda(res['preco_estimado'])} estimados"
    return moeda(res["preco"])


def _exclusoes(cat, opcao, res):
    return ct.textos_validos(cat, [it["servico"] for it in opcao["itens"]], bool(res["fornecimentos"]), "exclusoes")


def _premissas_fases(res):
    textos = []
    if len(res["fases"]) > 1:
        textos.append("As fases são sequenciais: cada fase se inicia após o aceite da anterior e pode ser contratada "
                      "separadamente, pelos valores indicados para ela.")
    for n, fr in enumerate(res["fases"], 1):
        nome = f"Fase {n} ({fr['nome']})" if len(res["fases"]) > 1 else fr["nome"]
        if fr["cobranca"] == "percentual":
            if fr["percentual"]["forma"] == "fixo_global":
                textos.append(f"{nome}: honorário fixo, calculado sobre o valor global estimado da obra; não se altera com "
                              "variações do valor da obra, salvo aditivo de escopo.")
            else:
                textos.append(f"{nome}: o valor das medições considerado é o aprovado pela fiscalização do contratante; "
                              "o honorário estimado acompanha o valor efetivamente medido.")
        elif fr["cobranca"] == "pacote":
            pk = fr["pacote"]
            if pk["forma"] == "mensal":
                saldo = "acumula para os meses seguintes, até o fim do contrato" if pk["acumula_saldo"] else \
                    "não acumula para os meses seguintes"
                textos.append(f"{nome}: pacote mensal de {pk['horas_mes']:.0f} horas, apontadas em relatório mensal de horas e "
                              f"atividades; o saldo não utilizado {saldo}.")
            else:
                textos.append(f"{nome}: pacote de {pk['horas']:.0f} horas, consumidas sob demanda em até {pk['meses']} meses "
                              "e apontadas em relatório mensal de horas e atividades.")
            textos.append(f"{nome}: horas excedentes, mediante autorização prévia, são cobradas a "
                          f"{moeda(pk['taxa_excedente'])} por hora.")
    return textos


def _premissas(cat, proposta, opcao, res):
    textos = [p for p in ct.textos_validos(cat, [it["servico"] for it in opcao["itens"]], bool(res["fornecimentos"]),
                                           "premissas") if p not in PREMISSAS_FINAIS]
    return textos + _premissas_fases(res) + PREMISSAS_FINAIS


def _clausulas_desempenho(n, fr, varias):
    d, cfg = fr["desempenho"], fr["_cfg_desempenho"]
    nome = f"da fase {n}" if varias else "da assessoria"
    criterios = []
    if cfg["economia_pct"] > 0:
        criterios.append(f"{pct_br(cfg['economia_pct'])} da economia apurada entre o valor global estimado da obra "
                         f"({moeda(d['base_global'])}) e o custo final")
    if cfg["prazo_pct_semana"] > 0:
        criterios.append(f"{pct_br(cfg['prazo_pct_semana'])} dos honorários {nome} por semana de antecipação da conclusão "
                         "em relação ao prazo previsto")
    if cfg["retrabalho_bonus_pct"] > 0:
        criterios.append(f"{pct_br(cfg['retrabalho_bonus_pct'])} dos honorários {nome} se o índice de retrabalho (custo de "
                         f"retrabalho ÷ valor medido) ficar igual ou abaixo de {pct_br(cfg['retrabalho_meta'])}")
    if not criterios:
        return []
    return [f"Taxa de desempenho {nome} (variável, adicional aos honorários): " + "; ".join(criterios) + ".",
            f"A taxa de desempenho é limitada a {pct_br(cfg['teto_pct'])} dos honorários {nome} ({moeda(d['teto'])}), apurada no "
            "encerramento da obra com base nas medições aprovadas e nos registros de prazo e de retrabalho, e paga na "
            "assinatura do termo de recebimento definitivo."]


def _pagamento(res):
    fases = res["fases"]
    varias = len(fases) > 1
    itens = []
    for n, fr in enumerate(fases, 1):
        pre = f"Fase {n} ({fr['nome']}): " if varias else ""
        for m in fr["pagamento"]:
            itens.append(f"{pre}{m['marco']}." if abs(m["pct"] - 1) < 1e-9 else f"{pre}{pct_br(m['pct'])} {m['marco']}.")
        if fr["desempenho"]:
            itens += _clausulas_desempenho(n, fr, varias)
    return [i[0].upper() + i[1:] for i in itens]


def _minuscula(texto):
    """Primeira letra minúscula, preservando siglas (TAB, PMOC, CFD...)."""
    if len(texto) > 1 and texto[1].islower():
        return texto[0].lower() + texto[1:]
    return texto


def _linhas_escopo(res):
    """Uma linha por serviço (numerada por fase quando há mais de uma) com atividades e entregáveis."""
    varias = len(res["fases"]) > 1
    saida = []
    for n, fr in enumerate(res["fases"], 1):
        grupos = []
        for l in fr["linhas"]:
            if not grupos or grupos[-1]["item"] != l["item"]:
                grupos.append({"item": l["item"], "nome": l["servico_nome"], "ativ": [], "entr": []})
            g = grupos[-1]
            g["ativ"].append(_minuscula(l["nome"]))
            if l["entregavel"] and l["entregavel"] not in g["entr"]:
                g["entr"].append(l["entregavel"])
        for k, g in enumerate(grupos, 1):
            num = f"{n}.{k}" if varias else f"{len(saida) + 1:02d}"
            saida.append((num, f"{g['nome']}: {'; '.join(g['ativ'])}.", "; ".join(g["entr"])))
        if not grupos and fr["cobranca"] == "pacote":
            num = f"{n}.1" if varias else f"{len(saida) + 1:02d}"
            saida.append((num, f"{fr['nome']}: fiscalização e planejamento da obra sob demanda, dentro do pacote de horas.",
                          "Relatórios mensais de horas e atividades"))
    return saida


def _linhas_prazos(res):
    fases = res["fases"]
    if len(fases) == 1:
        prazos = [(e["rotulo"], f"{e['dias_corridos']} dias", e["marco"]) for e in res["cronograma"]]
    else:
        prazos = []
        for n, fr in enumerate(fases, 1):
            etapas = [e for e in res["cronograma"] if e["fase"] == fr["id"]]
            meses = next((e["meses"] for e in etapas if e.get("meses")), None)
            dias = sum(e["dias_corridos"] for e in etapas)
            prazos.append((f"Fase {n} — {fr['nome']}", f"durante a obra ({meses} meses)" if meses else f"{dias} dias",
                           etapas[-1]["marco"] if etapas else "Conclusão da fase"))
    prazos.append(("Prazo total estimado", f"{res['prazo_dias_corridos']} dias", "Conclusão dos serviços"))
    return prazos


def _texto_objeto(proposta, opcao, res, varias):
    nomes = []
    for l in res["linhas"]:
        if l["servico_nome"] not in nomes:
            nomes.append(l["servico_nome"])
    base = (proposta.get("objeto") or "serviços técnicos de engenharia").rstrip(".")
    sufixo = f" — {opcao['nome']}" if varias else ""
    if len(res["fases"]) > 1:
        fases = "; ".join(f"Fase {n} — {_minuscula(fr['nome'])}" for n, fr in enumerate(res["fases"], 1))
        return (f"Prestação de serviços técnicos de engenharia, em {len(res['fases'])} fases: {base}{sufixo}. "
                f"Compreende: {fases}, conforme as atividades e os entregáveis relacionados na seção 3.")
    return (f"Prestação de serviços técnicos de engenharia: {base}{sufixo}. "
            f"Compreende: {'; '.join(nomes)}, conforme as atividades e os entregáveis relacionados na seção 3.")


# ----------------------------------------------------------------------------- API
def gerar_opcao(cat, proposta, opcao, res, varias=False):
    """Preenche o template para uma opção. Devolve (bytes, nome_arquivo, pendencias)."""
    params = proposta["parametros"]
    cfg_fases = {f["id"]: f for f in opcao["fases"]}
    for fr in res["fases"]:
        fr["_cfg_desempenho"] = cfg_fases.get(fr["id"], {}).get("desempenho", {})
    doc = Document(TEMPLATE)
    cl = proposta["cliente"]
    numero = proposta["numero"] + (f"-{opcao['id']}" if varias else "")
    campos = {
        "[Objeto — ex.: Laudo Técnico de Sistema de Climatização]": (proposta.get("objeto") or "Proposta de serviços de engenharia")
                                                                    + (f" — {opcao['nome']}" if varias else ""),
        "[PC-KTK-001/2026]": numero,
        "[__/__/____]": _data_br(proposta["data"]),
        "[15 dias corridos]": f"{proposta['validade_dias']} dias corridos",
        "[Razão social do contratante]": cl.get("razao") or "[Razão social do contratante]",
        "[Nome / cargo / e-mail / telefone]": cl.get("contato") or "[Nome / cargo / e-mail / telefone]",
        "[Logradouro, nº — Bairro — Cidade/UF]": cl.get("endereco") or "[Logradouro, nº — Bairro — Cidade/UF]",
        "[Razão social]": cl.get("razao") or "[Razão social]",
        "[Nome e cargo do signatário]": cl.get("signatario") or "[Nome e cargo do signatário]",
        "em __ dias da data": f"em {proposta.get('prazo_pagamento_dias', 15)} dias da data",
    }
    if any(d["descricao"] == "Deslocamento" for d in res["diretos"]):
        campos[FRASE_DESLOCAMENTO] = "Estão inclusos os deslocamentos previstos para as atividades de campo desta proposta."
    for par in list(_todos_paragrafos(doc)):
        for de, para in campos.items():
            substituir(par, de, para)
        if par.text.startswith("[Descreva o serviço contratado"):
            _definir(par, _texto_objeto(proposta, opcao, res, varias))

    rev = _tabela(doc, "Rev.")
    _definir(rev.rows[1].cells[0].paragraphs[0], proposta.get("revisao", "00"))
    _definir(rev.rows[1].cells[1].paragraphs[0], _data_br(proposta["data"]))

    _preencher_linhas(_tabela(doc, "#", "Serviço"), _linhas_escopo(res))
    _preencher_linhas(_tabela(doc, "Etapa", "Prazo"), _linhas_prazos(res))

    hon = _tabela(doc, "Descrição", "Qtd.")
    _preencher_linhas(hon, _linhas_honorarios(cat, params, opcao, res), manter_ultima=True)
    _definir(hon.rows[-1].cells[-1].paragraphs[0], _total_texto(res))

    secoes = _listas_por_secao(doc)
    _ajustar_lista(_lista(secoes, "EXCLUS"), _exclusoes(cat, opcao, res))
    _ajustar_lista(_lista(secoes, "PREMISSAS"), _premissas(cat, proposta, opcao, res))
    _ajustar_lista(_lista(secoes, "PAGAMENTO"), _pagamento(res))

    pendencias = sorted({m for p in _todos_paragrafos(doc) for m in re.findall(r"\[[^\]]{3,}\]", p.text)})
    buf = io.BytesIO()
    doc.save(buf)
    nome = f"KAPTEK-proposta-{slug(cl.get('razao'))}-{slug(proposta['numero'])}"
    nome += (f"-op{opcao['id']}" if varias else "") + f"-rev{proposta.get('revisao', '00')}.docx"
    return buf.getvalue(), nome, pendencias


def gerar(cat, proposta, resultado):
    """Uma proposta por opção ofertada (a estimativa interna não vira proposta): lista de (bytes, nome, pendencias)."""
    pares = [(o, r) for o, r in zip(proposta["opcoes"], resultado["opcoes"]) if not r["interno"]]
    varias = len(pares) > 1
    return [gerar_opcao(cat, proposta, o, r, varias) for o, r in pares]


def verificar(conteudo):
    """Verificação da habilidade KAPTEK: margens 3,2/3,0 e logo no cabeçalho."""
    d = Document(io.BytesIO(conteudo))
    s = d.sections[0]
    return {"margem_sup_cm": round(s.top_margin.cm, 1), "margem_esq_cm": round(s.left_margin.cm, 1),
            "logo": "graphicData" in s.header._element.xml}
