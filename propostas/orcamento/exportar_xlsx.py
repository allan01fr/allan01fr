"""
Planilha de escopo e preços, no padrão visual KAPTEK, com fórmulas vivas.

Toda conta do motor é refeita em fórmula: mudar uma produtividade, um custo/h,
a margem ou uma quantidade na planilha recalcula horas, preço, prazo e parcelas.
Só usa funções compatíveis com o Excel 2007 (SUM, SUMIF, SUMIFS, INDEX/MATCH, ROUNDUP, IF, MIN).
Cada opção tem uma matriz de preço com uma coluna por fase (escopo fechado, hora técnica, % da obra,
pacote de horas e taxa de desempenho) e a coluna Total.

Convenção de cores: fundo amarelo e texto azul = entrada; texto verde = vínculo
com outra aba; grafite = fórmula.
"""

import io
import json
import math
import os

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

from . import motor
from . import parametros as pm

RAIZ = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
LOGO = os.path.join(RAIZ, "modelos", "logo-horizontal.png")
ORG = os.path.join(RAIZ, "dados", "kaptek_org.json")

AZUL, GRAFITE, COBRE, CINZA = "1B3A5C", "3E4448", "C18E6D", "8A8A8A"
FUNDO_CAB, FILETE, AMARELO, AZUL_ENTRADA, VERDE = "F2F4F6", "D8D8D8", "FFF2CC", "0000FF", "008000"
FONTE = "Arial"

MOEDA = '"R$" #,##0.00;("R$" #,##0.00);"-"'
HORAS = '#,##0.0;(#,##0.0);"-"'
NUM = '#,##0.00;(#,##0.00);"-"'
PCT = '0.0%;(0.0%);"-"'
INT = '0;(0);"-"'

_fino = Side(style="thin", color=FILETE)
BORDA = Border(left=_fino, right=_fino, top=_fino, bottom=_fino)


def _fill(cor):
    return PatternFill("solid", start_color=cor, end_color=cor)


class Folha:
    """Atalhos de escrita com o estilo da marca."""

    def __init__(self, ws):
        self.ws = ws

    def c(self, linha, col, valor=None, fmt=None, negrito=False, cor=GRAFITE, fundo=None, alinhar=None,
          quebra=False, tam=10, borda=True):
        cel = self.ws.cell(linha, col, valor)
        cel.font = Font(name=FONTE, size=tam, bold=negrito, color=cor)
        if fmt:
            cel.number_format = fmt
        if fundo:
            cel.fill = _fill(fundo)
        if alinhar or quebra:
            cel.alignment = Alignment(horizontal=alinhar, vertical="center", wrap_text=quebra)
        if borda:
            cel.border = BORDA
        return cel

    def entrada(self, linha, col, valor, fmt=None):
        return self.c(linha, col, valor, fmt, cor=AZUL_ENTRADA, fundo=AMARELO)

    def vinculo(self, linha, col, formula, fmt=None, negrito=False):
        return self.c(linha, col, formula, fmt, negrito=negrito, cor=VERDE)

    def titulo(self, linha, texto, ate_col):
        self.ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=ate_col)
        cel = self.c(linha, 1, texto, negrito=True, cor="FFFFFF", fundo=AZUL, tam=13, borda=False)
        cel.alignment = Alignment(vertical="center", indent=1)
        self.ws.row_dimensions[linha].height = 26
        cobre = Side(style="medium", color=COBRE)
        for col in range(1, ate_col + 1):
            self.ws.cell(linha, col).border = Border(bottom=cobre)

    def secao(self, linha, texto, ate_col):
        self.ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=ate_col)
        cel = self.c(linha, 1, texto, negrito=True, cor=AZUL, tam=11, borda=False)
        cel.border = Border(bottom=Side(style="thin", color=COBRE))
        self.ws.row_dimensions[linha].height = 20

    def cabecalho(self, linha, rotulos, col=1):
        for i, r in enumerate(rotulos):
            self.c(linha, col + i, r, negrito=True, cor=AZUL, fundo=FUNDO_CAB, alinhar="center", quebra=True)
        self.ws.row_dimensions[linha].height = 30

    def nota(self, linha, texto, ate_col, col=1):
        self.ws.merge_cells(start_row=linha, start_column=col, end_row=linha, end_column=ate_col)
        cel = self.c(linha, col, texto, cor=CINZA, tam=8, borda=False)
        cel.alignment = Alignment(wrap_text=True, vertical="top")

    def larguras(self, mapa):
        for col, w in mapa.items():
            self.ws.column_dimensions[col].width = w


def _q(nome):
    """Nome de aba entre aspas para uso em fórmula."""
    return "'" + nome.replace("'", "''") + "'"


def _org():
    try:
        with open(ORG, encoding="utf-8") as f:
            return json.load(f)
    except OSError:
        return {}


def _logo(ws, ancora):
    if os.path.exists(LOGO):
        img = XLImage(LOGO)
        img.width, img.height = 200, round(200 * img.height / img.width)
        ws.add_image(img, ancora)


def _impressao(ws, paisagem=True):
    ws.page_setup.orientation = "landscape" if paisagem else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True


# ----------------------------------------------------------------------------- Parâmetros
def _aba_parametros(wb, cat, proposta):
    params = proposta["parametros"]
    ws = wb.create_sheet("Parametros")
    f = Folha(ws)
    f.larguras({"A": 10, "B": 40, "C": 18, "D": 18, "E": 60})
    f.titulo(1, "PARÂMETROS COMERCIAIS DESTA PROPOSTA", 5)
    ref = {}

    f.secao(3, "Custo e taxa de venda por perfil", 5)
    f.cabecalho(4, ["Código", "Perfil", "Custo/h (R$)", "Taxa de venda/h (R$)", "Observação"])
    ref["custo"], ref["taxa"] = {}, {}
    for i, (cod, nome) in enumerate(cat["perfis"].items()):
        r = 5 + i
        f.c(r, 1, cod, alinhar="center")
        f.c(r, 2, nome)
        f.entrada(r, 3, float(params["perfis"][cod]["custo_h"]), MOEDA)
        f.entrada(r, 4, float(params["perfis"][cod]["taxa_venda_h"]), MOEDA)
        f.c(r, 5, "Salário + encargos + benefícios ÷ horas produtivas" if i == 0 else None, cor=CINZA, tam=8)
        ref["custo"][cod] = f"Parametros!$C${r}"
        ref["taxa"][cod] = f"Parametros!$D${r}"
    r = 5 + len(cat["perfis"]) + 1

    f.secao(r, "Formação do preço", 5)
    r += 1
    imp = pm.aliquota_impostos(params)
    regime = params["impostos"]["regime"]
    nota_imp = ("Simples Nacional — alíquota efetiva informada" if regime == "simples"
                else "Lucro Presumido: PIS + COFINS + IRPJ/CSLL presumidos + ISS")
    linhas = [
        ("overhead_metodo", "Método de overhead (percentual | hora)", params["overhead"]["metodo"], None, "percentual = % sobre a MO; hora = custo fixo ÷ horas faturáveis × horas da proposta"),
        ("overhead_pct", "Overhead sobre a mão de obra", float(params["overhead"]["pct"]), PCT, "Estrutura, aluguel, administrativo, software"),
        ("custo_fixo", "Custo fixo mensal do escritório", float(params["overhead"]["custo_fixo_mensal"]), MOEDA, "Usado no método hora"),
        ("horas_fat", "Horas faturáveis por mês", float(params["overhead"]["horas_faturaveis_mes"]), NUM, "Usado no método hora"),
        ("imp", "Impostos sobre o faturamento", imp, PCT, nota_imp),
        ("margem", "Margem-alvo sobre o preço", float(params["margem"]), PCT, "Lucro desejado"),
        ("cont", f"Contingência (risco {proposta['contexto']['risco']})", float(params["contingencia"][proposta["contexto"]["risco"]]), PCT, "Reserva de risco do escopo fechado"),
        ("adm", "Administração sobre reembolsáveis (hora técnica)", float(params["reembolsaveis_adm"]), PCT, "Aplicada aos custos diretos no modelo hora técnica"),
        ("horas_dia", "Horas por dia de trabalho", float(params["cronograma"]["horas_dia"]), NUM, "Base do cronograma"),
        ("area", "Área de referência", float(proposta["contexto"].get("area_ref") or 0), NUM, "m² — usada nos indicadores R$/m² e CUB (0 = não calcular)"),
        ("cub_valor", f"CUB/m² ({params['cub'].get('padrao', '')} {params['cub'].get('mes', '')})".strip(), float(params["cub"].get("valor") or 0), MOEDA,
         f"{params['cub'].get('fonte', 'SINDUSCON')} — 0 = sem referência de CUB"),
        ("cub_pct", "% do CUB × área para honorários de referência", float(params["cub"].get("pct_honorarios") or 0), PCT, "Faixa usual 2% a 11%"),
        ("garantia_pct", f"Provisão de garantia — turnkey ({params['turnkey']['garantia_meses']} meses)", float(params["turnkey"]["garantia_pct"]), PCT,
         "Sobre o custo de fornecimento e execução"),
        ("cesta_faixa", "Cesta de preços: faixa aceita em torno da mediana", float(params["cesta"]["faixa_alerta"]), PCT,
         "Referências fora da faixa ficam marcadas e saem do valor adotado (aba Pesquisa de preços)"),
    ]
    for chave, rot, val, fmt, obs in linhas:
        f.c(r, 2, rot)
        f.entrada(r, 3, val, fmt)
        f.c(r, 5, obs, cor=CINZA, tam=8)
        ref[chave] = f"Parametros!$C${r}"
        r += 1
    f.c(r, 2, "Overhead por hora (método hora)", negrito=True)
    f.c(r, 3, f"=IF({ref['horas_fat']}>0,{ref['custo_fixo']}/{ref['horas_fat']},0)", MOEDA, fundo=FUNDO_CAB)
    ref["overhead_h"] = f"Parametros!$C${r}"
    r += 2

    f.secao(r, "BDI de fornecimento e execução (turnkey)", 5)
    r += 1
    f.cabecalho(r, ["Código", "Categoria", "BDI", "", "Observação"])
    r += 1
    rot_cat = {"material": "Material", "equipamento": "Equipamento", "mo_execucao": "Mão de obra de execução", "subcontrato": "Subcontrato"}
    ini = r
    for k, v in params["bdi_fornecimento"].items():
        f.c(r, 1, k, alinhar="center")
        f.c(r, 2, rot_cat.get(k, k))
        f.entrada(r, 3, float(v), PCT)
        r += 1
    ref["bdi_chaves"] = f"Parametros!$A${ini}:$A${r - 1}"
    ref["bdi_valores"] = f"Parametros!$C${ini}:$C${r - 1}"
    f.nota(r, params.get("notas", {}).get("bdi_fornecimento", ""), 5)
    r += 2
    for k in ("impostos", "custos", "art"):
        if params.get("notas", {}).get(k):
            f.nota(r, params["notas"][k], 5)
            r += 1
    ws.freeze_panes = "A5"
    _impressao(ws, paisagem=False)
    return ref


# ----------------------------------------------------------------------------- Opção
COL_PERFIL = 10     # J: primeira coluna de horas por perfil (A fase, B #, C serviço ... I fator)
ROT_COBRANCA = motor.ROTULO_COBRANCA


def _mix_formula(params, ref, chave):
    """Média ponderada pelo mix padrão do pacote (sem serviços na fase), em fórmula."""
    mix = params["cobranca"]["pacote"]["mix_padrao"]
    soma = sum(mix.values()) or 1
    termos = "+".join(f"{v}*{ref[chave][p]}" for p, v in mix.items() if p in ref[chave])
    return f"(({termos})/{soma})"


def _aba_opcao(wb, cat, proposta, opcao, res, ref, celulas):
    params = proposta["parametros"]
    perfis = list(cat["perfis"])
    n_p = len(perfis)
    nome = f"Opção {opcao['id']}"[:31]
    ws = wb.create_sheet(nome)
    f = Folha(ws)
    c_tot = COL_PERFIL + n_p          # total h
    c_mo = c_tot + 1                  # custo MO
    c_fixo = c_mo + 2                 # bloco de produtividade (deixa uma coluna de folga)
    c_unid = c_fixo + n_p
    ultima = c_unid + n_p - 1
    interno = res["interno"]
    fases = res["fases"]
    cfg = {x["id"]: x for x in opcao["fases"]}
    f.titulo(1, f"{opcao['nome'].upper()} — " + ("ESTIMATIVA INTERNA DO VALOR GLOBAL (NÃO OFERTADA)" if interno
                                                 else "FASES, HORAS E PREÇO"), c_mo)
    f.nota(2, "Horas por perfil = (fixo + unid × qtd^b) × fator. Edite as células amarelas: quantidade, expoente, "
              "fator e a produtividade (fixo/unid) à direita. A coluna Fase liga cada atividade ao preço da sua fase.", c_mo)

    cab = ["Fase", "#", "Serviço", "Atividade", "Entregável", "Etapa", "Qtd (driver)", "b", "Fator"]
    cab += [f"h {p}" for p in perfis] + ["Total (h)", "Custo MO (R$)", ""]
    cab += [f"Fixo {p}" for p in perfis] + [f"Unid {p}" for p in perfis]
    f.cabecalho(3, cab)
    ws.cell(3, c_mo + 1).fill = PatternFill()
    rot_etapa = {e["id"]: e["rotulo"] for e in cat["etapas"]}

    r0 = r = 4
    for fr in fases:
        for l in fr["linhas"]:
            f.c(r, 1, fr["id"], alinhar="center")
            f.c(r, 2, l["item"] + 1, INT, alinhar="center")
            f.c(r, 3, l["servico_nome"], quebra=True, tam=9)
            f.c(r, 4, l["nome"], quebra=True, tam=9)
            f.c(r, 5, l["entregavel"], quebra=True, tam=9)
            f.c(r, 6, rot_etapa[l["etapa"]], quebra=True, tam=9)
            f.entrada(r, 7, round(l["q"], 6), NUM)
            f.entrada(r, 8, l["b"], NUM)
            f.entrada(r, 9, round(l["fator"], 6), "0.000")
            for k, p in enumerate(perfis):
                cf, cu = L(c_fixo + k), L(c_unid + k)
                f.entrada(r, c_fixo + k, l["fixo"][p], HORAS)
                f.entrada(r, c_unid + k, l["unid"][p], "0.0000")
                f.c(r, COL_PERFIL + k, f"=({cf}{r}+{cu}{r}*$G{r}^$H{r})*$I{r}", HORAS)
            f.c(r, c_tot, f"=SUM({L(COL_PERFIL)}{r}:{L(COL_PERFIL + n_p - 1)}{r})", HORAS, negrito=True)
            mo = "+".join(f"{L(COL_PERFIL + k)}{r}*{ref['custo'][p]}" for k, p in enumerate(perfis))
            f.c(r, c_mo, "=" + mo, MOEDA)
            ws.row_dimensions[r].height = 30
            r += 1
    r_fim = max(r - 1, r0)
    rt = r_fim + 1
    f.c(rt, 3, "TOTAL", negrito=True, fundo=FUNDO_CAB)
    for col in (1, 2, *range(4, COL_PERFIL)):
        f.c(rt, col, None, fundo=FUNDO_CAB)
    for col in range(COL_PERFIL, c_mo + 1):
        letra = L(col)
        f.c(rt, col, f"=SUM({letra}{r0}:{letra}{r_fim})", MOEDA if col == c_mo else HORAS, negrito=True, fundo=FUNDO_CAB)
    FA = f"$A${r0}:$A${r_fim}"
    col_tot = f"${L(c_tot)}${r0}:${L(c_tot)}${r_fim}"
    col_mo = f"${L(c_mo)}${r0}:${L(c_mo)}${r_fim}"
    col_h = {p: f"${L(COL_PERFIL + k)}${r0}:${L(COL_PERFIL + k)}${r_fim}" for k, p in enumerate(perfis)}

    # custos diretos (por fase: cada fase tem a sua ART e o seu deslocamento)
    r = rt + 2
    f.secao(r, "Custos diretos (instrumentos, terceiros, deslocamento, ART, garantia)", 8)
    r += 1
    f.cabecalho(r, ["Fase", "Descrição", "Serviço / memória", "Unid", "Qtd", "Valor unit. (R$)", "Valor (R$)"])
    r += 1
    d0 = r
    for fr in fases:
        for d in fr["diretos"]:
            f.c(r, 1, fr["id"], alinhar="center")
            f.c(r, 2, d["descricao"], quebra=True, tam=9)
            f.c(r, 3, d["servico_nome"], quebra=True, tam=9)
            f.c(r, 4, d["unid"], alinhar="center")
            if d.get("garantia"):
                f.vinculo(r, 5, f"=SUMIF(Fornecimentos!$A:$A,\"{opcao['id']}\",Fornecimentos!$G:$G)", MOEDA)
                f.vinculo(r, 6, f"={ref['garantia_pct']}", PCT)
            else:
                f.entrada(r, 5, round(d["qtd"], 6), NUM)
                f.entrada(r, 6, round(d["valor_unit"], 2), MOEDA)
            f.c(r, 7, f"=E{r}*F{r}", MOEDA)
            r += 1
    d_fim = max(r - 1, d0)
    if r == d0:
        r += 1
    f.c(r, 2, "TOTAL DE CUSTOS DIRETOS", negrito=True, fundo=FUNDO_CAB)
    f.c(r, 7, f"=SUM(G{d0}:G{d_fim})", MOEDA, negrito=True, fundo=FUNDO_CAB)
    FD, col_dir = f"$A${d0}:$A${d_fim}", f"$G${d0}:$G${d_fim}"

    # matriz de preço: uma coluna por fase + total
    nf = len(fases)
    col_f = [3 + k for k in range(nf)]
    c_total = 3 + nf
    c_obs = c_total + 1
    algum_desemp = any(fr["desempenho"] for fr in fases)
    r += 2
    f.secao(r, "Preço por fase" + (" — estimativa interna" if interno else ""), c_obs)
    r += 1
    f.cabecalho(r, [fr["nome"] for fr in fases] + ["Total", "Observação"], col=3)
    f.c(r, 2, "Fase", negrito=True, cor=AZUL, fundo=FUNDO_CAB)
    r += 1
    SECAO = None
    linhas = [
        ("cob", "Cobrança (fechado | hora | percentual | pacote)", None, False, "Entrada"),
        (SECAO, "Custo e escopo fechado"),
        ("h_serv", "Horas dos serviços da fase", HORAS, True, "Atividades da fase (tabela acima)"),
        ("mo", "Custo de mão de obra", MOEDA, True, "Horas × custo/h de cada perfil"),
        ("overhead", "Overhead", MOEDA, True, "Conforme o método da aba Parametros"),
        ("dir", "Custos diretos", MOEDA, True, "Inclui ART e deslocamento da fase"),
        ("custo", "CUSTO DOS SERVIÇOS", MOEDA, True, "MO + overhead + diretos"),
        ("cont", "Contingência", MOEDA, True, "Custo × contingência"),
        ("preco_serv", "PREÇO DOS SERVIÇOS", MOEDA, True, "(Custo + contingência) ÷ (1 − impostos − margem)"),
        ("impostos", "  Impostos", MOEDA, True, "Preço × alíquota"),
        ("margem", "  Margem", MOEDA, True, "Preço × margem-alvo"),
        ("custo_forn", "Custo dos fornecimentos", MOEDA, True, "Aba Fornecimentos"),
        ("preco_forn", "Preço dos fornecimentos (com BDI)", MOEDA, True, "Custo × (1 + BDI da categoria)"),
        ("fechado", "PREÇO — ESCOPO FECHADO", MOEDA, True, "Serviços + fornecimentos"),
        (SECAO, "Hora técnica"),
        ("honor", "Honorários (horas × taxa de venda)", MOEDA, True, "Taxas por perfil na aba Parametros"),
        ("reemb", "Reembolsáveis", MOEDA, True, "Diretos × (1 + administração) ÷ (1 − impostos)"),
        ("hora", "PREÇO — HORA TÉCNICA (estimado)", MOEDA, True, "Honorários + reembolsáveis + fornecimentos"),
        ("hora_teto", "Hora técnica — teto", MOEDA, True, "Horas com a contingência"),
        (SECAO, "% da obra"),
        ("pct", "% sobre a obra", PCT, False, "Entrada"),
        ("forma", "Forma (medicoes | fixo_global | fixo_mais_medicoes)", None, False, "Entrada"),
        ("fixa", "Parcela fixa (forma fixo + medições)", MOEDA, False, "Entrada"),
        ("base", "Valor global da obra (base)", MOEDA, False, "Verde = estimativa interna; amarelo = informado"),
        ("meses", "Meses de obra (0 = prazo pelas horas)", INT, False, "Entrada"),
        ("pct_fixo", "Honorário % — parte fixa", MOEDA, True, "Fixo pelo global, ou a parcela fixa"),
        ("pct_est", "Honorário % — sobre as medições (estimado)", MOEDA, True, "% × base"),
        ("pct_hon", "HONORÁRIO — % DA OBRA", MOEDA, True, "Parte fixa + estimada"),
        ("pct_min", "% mínimo para cobrir o custo (margem zero)", PCT, False, "Custo ÷ (1 − impostos) ÷ base"),
        (SECAO, "Pacote de horas"),
        ("pk_horas", "Horas do pacote", HORAS, True, "Entrada; com fórmula = horas dos serviços da fase"),
        ("pk_cheia", "Taxa média cheia (mix de perfis)", MOEDA, False, "Mix dos serviços da fase ou mix padrão"),
        ("pk_desc", "Desconto do pacote", PCT, False, "Entrada"),
        ("pk_taxa", "Taxa do pacote (R$/h)", MOEDA, False, "Taxa cheia × (1 − desconto)"),
        ("pk_hon", "Honorários do pacote", MOEDA, True, "Horas × taxa"),
        ("pk_custo", "Custo das horas do pacote", MOEDA, True, "Horas × custo médio + overhead + diretos"),
        ("pk_parc", "Parcelas", INT, False, "Mensal: meses; total: parcelas"),
        ("pk_valor_parc", "Valor da parcela", MOEDA, False, "(Honorários + reembolsáveis) ÷ parcelas"),
        ("pk_exc", "Hora excedente (R$/h)", MOEDA, False, "Taxa cheia × (1 + acréscimo)"),
        (SECAO, "Preço da fase"),
        ("honorarios", "Honorários da fase (conforme a cobrança)", MOEDA, True, ""),
        ("custo_fase", "Custo da fase", MOEDA, True, "No pacote, o custo das horas vendidas"),
        ("p_fixo", "Preço — parcela fixa", MOEDA, True, "Escopo fechado, pacote, % fixo e parcela fixa"),
        ("p_est", "Preço — parcela estimada", MOEDA, True, "Hora técnica e % sobre as medições"),
        ("preco", "PREÇO DA FASE", MOEDA, True, "Fixo + estimado"),
        ("h_fase", "Horas da fase", HORAS, True, "No pacote, as horas vendidas"),
        ("minimo", "Preço mínimo (margem zero)", MOEDA, True, "Custo ÷ (1 − impostos)"),
        ("me", "Margem efetiva", PCT, False, "(Honorários × (1 − impostos) − custo) ÷ honorários"),
    ]
    if algum_desemp:
        linhas += [
            (SECAO, "Taxa de desempenho (variável, fora do preço)"),
            ("de_part", "Participação na economia sobre o valor global", PCT, False, "Entrada"),
            ("de_cen", "Cenário: economia obtida", PCT, False, "Entrada (simulação)"),
            ("de_sem", "% dos honorários por semana antecipada", PCT, False, "Entrada"),
            ("de_dias", "Cenário: dias antecipados", INT, False, "Entrada (simulação)"),
            ("de_bonus", "Bônus por retrabalho dentro da meta", PCT, False, "Entrada"),
            ("de_ok", "Cenário: meta de retrabalho atingida (1 = sim)", INT, False, "Entrada (simulação)"),
            ("de_teto", "Teto (% dos honorários)", PCT, False, "Entrada"),
            ("de_eco", "Bônus — economia", MOEDA, True, "Participação × economia × base"),
            ("de_prazo", "Bônus — prazo", MOEDA, True, "% × honorários × semanas"),
            ("de_ret", "Bônus — retrabalho", MOEDA, True, "% × honorários, se a meta for atingida"),
            ("de_teto_v", "Teto da taxa de desempenho", MOEDA, True, "% × honorários"),
            ("de_pot", "TAXA DE DESEMPENHO — POTENCIAL", MOEDA, True, "Soma dos bônus, limitada ao teto"),
        ]
    linhas += [
        (SECAO, "Indicadores da opção"),
        ("rs_h", "Preço dos serviços por hora (R$/h)", MOEDA, None, "Preço dos serviços ÷ horas"),
        ("rs_m2", "Preço por m² (R$/m²)", MOEDA, None, "Preço ÷ área de referência"),
        ("ref_cub", "Referência: % do CUB × área", MOEDA, None, "Comparação com honorários usuais (aba Parametros)"),
    ]
    pos = {}
    for item in linhas:
        if item[0] is SECAO:
            f.c(r, 2, item[1], negrito=True, cor=AZUL)
            r += 1
            continue
        pos[item[0]] = r
        r += 1

    imp, mg, cont, adm = ref["imp"], ref["margem"], ref["cont"], ref["adm"]
    ovm, ovp, ovh = ref["overhead_metodo"], ref["overhead_pct"], ref["overhead_h"]
    taxa_mix, custo_mix = _mix_formula(params, ref, "taxa"), _mix_formula(params, ref, "custo")
    oid = opcao["id"]
    fq = "Fornecimentos"
    for k, fr in enumerate(fases):
        c = L(col_f[k])
        X = {ch: f"{c}{v}" for ch, v in pos.items()}
        fc = cfg.get(fr["id"], opcao["fases"][0])
        fid = fr["id"]
        if interno:   # a estimativa soma todos os fornecimentos da opção
            filtro_f = f"{fq}!$A:$A,\"{oid}\""
            custo_forn, preco_forn = f"=SUMIF({filtro_f},{fq}!$G:$G)", f"=SUMIF({filtro_f},{fq}!$I:$I)"
        else:
            filtro_f = f"{fq}!$A:$A,\"{oid}\",{fq}!$M:$M,\"{fid}\""
            custo_forn, preco_forn = f"=SUMIFS({fq}!$G:$G,{filtro_f})", f"=SUMIFS({fq}!$I:$I,{filtro_f})"
        honor = "+".join(f"SUMIF({FA},\"{fid}\",{col_h[p]})*{ref['taxa'][p]}" for p in perfis)
        ob, pc, pk = fc["obra"], fc["percentual"], fc["pacote"]
        cm = f"IF({X['h_serv']}>0,{X['mo']}/{X['h_serv']},{custo_mix})"
        formulas = {
            "h_serv": f"=SUMIF({FA},\"{fid}\",{col_tot})",
            "mo": f"=SUMIF({FA},\"{fid}\",{col_mo})",
            "overhead": f"=IF({ovm}=\"hora\",{X['h_serv']}*{ovh},{X['mo']}*{ovp})",
            "dir": f"=SUMIF({FD},\"{fid}\",{col_dir})",
            "custo": f"={X['mo']}+{X['overhead']}+{X['dir']}",
            "cont": f"={X['custo']}*{cont}",
            "preco_serv": f"=({X['custo']}+{X['cont']})/(1-{imp}-{mg})",
            "impostos": f"={X['preco_serv']}*{imp}",
            "margem": f"={X['preco_serv']}*{mg}",
            "custo_forn": custo_forn,
            "preco_forn": preco_forn,
            "fechado": f"={X['preco_serv']}+{X['preco_forn']}",
            "honor": "=" + honor,
            "reemb": f"={X['dir']}*(1+{adm})/(1-{imp})",
            "hora": f"={X['honor']}+{X['reemb']}+{X['preco_forn']}",
            "hora_teto": f"={X['honor']}*(1+{cont})+{X['reemb']}+{X['preco_forn']}",
            "pct_fixo": f"=IF({X['forma']}=\"fixo_global\",{X['pct']}*{X['base']},IF({X['forma']}=\"fixo_mais_medicoes\",{X['fixa']},0))",
            "pct_est": f"=IF({X['forma']}=\"fixo_global\",0,{X['pct']}*{X['base']})",
            "pct_hon": f"={X['pct_fixo']}+{X['pct_est']}",
            "pct_min": f"=IF({X['base']}>0,{X['custo']}/(1-{imp})/{X['base']},0)",
            "pk_cheia": f"=IF({X['h_serv']}>0,{X['honor']}/{X['h_serv']},{taxa_mix})",
            "pk_taxa": f"={X['pk_cheia']}*(1-{X['pk_desc']})",
            "pk_hon": f"={X['pk_horas']}*{X['pk_taxa']}",
            "pk_custo": f"={X['pk_horas']}*{cm}+IF({ovm}=\"hora\",{X['pk_horas']}*{ovh},{X['pk_horas']}*{cm}*{ovp})+{X['dir']}",
            "pk_valor_parc": f"=IF({X['pk_parc']}>0,({X['pk_hon']}+{X['reemb']})/{X['pk_parc']},0)",
            "pk_exc": f"={X['pk_cheia']}*(1+{float(pk.get('excedente_pct') or 0)})",
            "honorarios": (f"=IF({X['cob']}=\"hora\",{X['honor']}+{X['reemb']},IF({X['cob']}=\"percentual\",{X['pct_hon']},"
                           f"IF({X['cob']}=\"pacote\",{X['pk_hon']}+{X['reemb']},{X['preco_serv']})))"),
            "custo_fase": f"=IF({X['cob']}=\"pacote\",{X['pk_custo']},{X['custo']})",
            "p_fixo": (f"=IF({X['cob']}=\"fechado\",{X['fechado']},IF({X['cob']}=\"percentual\",{X['pct_fixo']}+{X['preco_forn']},"
                       f"IF({X['cob']}=\"pacote\",{X['pk_hon']}+{X['reemb']}+{X['preco_forn']},0)))"),
            "p_est": f"=IF({X['cob']}=\"hora\",{X['hora']},IF({X['cob']}=\"percentual\",{X['pct_est']},0))",
            "preco": f"={X['p_fixo']}+{X['p_est']}",
            "h_fase": f"=IF({X['cob']}=\"pacote\",{X['pk_horas']},{X['h_serv']})",
            "minimo": f"={X['custo_fase']}/(1-{imp})",
            "me": f"=IF({X['honorarios']}>0,({X['honorarios']}*(1-{imp})-{X['custo_fase']})/{X['honorarios']},0)",
        }
        entradas = {"cob": (fr["cobranca"], None), "pct": (float(pc["pct"]), PCT), "forma": (pc.get("forma", "medicoes"), None),
                    "fixa": (float(pc.get("parcela_fixa") or 0), MOEDA), "meses": (int(ob.get("meses") or 0), INT),
                    "pk_desc": (float(pk.get("desconto") or 0), PCT)}
        if fr["cobranca"] == "pacote":
            entradas["pk_parc"] = (int(fr["pacote"]["parcelas"]), INT)
        else:
            entradas["pk_parc"] = (int(pk.get("meses") if pk.get("forma", "mensal") == "mensal" else pk.get("parcelas") or 1), INT)
        horas_inf = float(pk.get("horas_mes") or 0) * int(pk.get("meses") or 1) if pk.get("forma", "mensal") == "mensal" \
            else float(pk.get("horas") or 0)
        if horas_inf > 0:
            entradas["pk_horas"] = (horas_inf, HORAS)
        else:
            formulas["pk_horas"] = f"={X['h_serv']}"
        # base da obra: vínculo com a estimativa interna, custo dos fornecimentos da fase ou valor informado
        est = celulas.get(ob.get("base_opcao")) if ob.get("base_modo") == "opcao" else None
        if est and est.get("interno"):
            formulas["base"] = "=" + (est["global"] if ob.get("base_parte") == "global" else est["obra"])
        elif ob.get("base_modo") == "opcao":
            formulas["base"] = "=0"
        elif ob.get("base_modo") == "fornecimentos":
            formulas["base"] = f"={X['custo_forn']}"
        else:
            entradas["base"] = (float(ob.get("base_valor") or 0), MOEDA)
        if algum_desemp:
            if fr["desempenho"]:
                de = fc["desempenho"]
                cen = de.get("cenario", {})
                entradas.update({"de_part": (float(de["economia_pct"]), PCT), "de_cen": (float(cen.get("economia", 0)), PCT),
                                 "de_sem": (float(de["prazo_pct_semana"]), PCT), "de_dias": (int(cen.get("dias_antecipados", 0)), INT),
                                 "de_bonus": (float(de["retrabalho_bonus_pct"]), PCT),
                                 "de_ok": (1 if cen.get("retrabalho_ok") else 0, INT), "de_teto": (float(de["teto_pct"]), PCT)})
                formulas.update({
                    "de_eco": f"={X['de_part']}*{X['de_cen']}*{X['base']}",
                    "de_prazo": f"={X['de_sem']}*{X['honorarios']}*{X['de_dias']}/7",
                    "de_ret": f"={X['de_bonus']}*{X['honorarios']}*{X['de_ok']}",
                    "de_teto_v": f"={X['de_teto']}*{X['honorarios']}",
                    "de_pot": f"=MIN({X['de_eco']}+{X['de_prazo']}+{X['de_ret']},{X['de_teto_v']})",
                })
            else:
                for ch in ("de_part", "de_cen", "de_sem", "de_dias", "de_bonus", "de_ok", "de_teto",
                           "de_eco", "de_prazo", "de_ret", "de_teto_v", "de_pot"):
                    formulas[ch] = "—"
        for item in linhas:
            if item[0] is SECAO or item[3] is None:
                continue
            ch, rot, fmt = item[0], item[1], item[2]
            forte = rot.isupper()
            if ch in entradas:
                val, fmt_e = entradas[ch]
                f.entrada(pos[ch], col_f[k], val, fmt_e or fmt)
            elif ch == "base" and formulas["base"].startswith("=") and est:
                f.vinculo(pos[ch], col_f[k], formulas[ch], fmt)
            else:
                f.c(pos[ch], col_f[k], formulas[ch], fmt, negrito=forte, fundo=FUNDO_CAB if forte else None,
                    alinhar="right" if formulas[ch] == "—" else None)

    T = L(c_total)
    TT = {ch: f"{T}{v}" for ch, v in pos.items()}
    ini, fim = L(col_f[0]), L(col_f[-1])
    for item in linhas:
        if item[0] is SECAO:
            continue
        ch, rot, fmt, soma, obs = item
        forte = rot.isupper()
        f.c(pos[ch], 2, rot, negrito=forte, fundo=FUNDO_CAB if forte else None)
        f.c(pos[ch], c_obs, obs, cor=CINZA, tam=8, fundo=FUNDO_CAB if forte else None)
        if soma:
            f.c(pos[ch], c_total, f"=SUM({ini}{pos[ch]}:{fim}{pos[ch]})", fmt, negrito=True, fundo=FUNDO_CAB)
        elif soma is False:
            f.c(pos[ch], c_total, None, fundo=FUNDO_CAB)
    f.c(pos["me"], c_total, f"=IF({TT['honorarios']}>0,({TT['honorarios']}*(1-{imp})-{TT['custo_fase']})/{TT['honorarios']},0)",
        PCT, negrito=True, fundo=FUNDO_CAB)
    f.c(pos["rs_h"], c_total, f"=IF({TT['h_fase']}>0,{TT['preco_serv']}/{TT['h_fase']},0)", MOEDA, negrito=True)
    f.c(pos["rs_m2"], c_total, f"=IF({ref['area']}>0,{TT['preco']}/{ref['area']},0)", MOEDA, negrito=True)
    f.c(pos["ref_cub"], c_total, f"={ref['cub_pct']}*{ref['cub_valor']}*{ref['area']}", MOEDA, negrito=True)

    if res.get("checklist_turnkey"):
        r += 1
        f.secao(r, "Etapas da obra consideradas no valor global", c_obs)
        r += 1
        rot_status = {"atendida": "Incluída", "dispensada": "Fora da estimativa", "opcional": "Opcional — não incluída",
                      "faltando": "FALTANDO"}
        for e in res["checklist_turnkey"]:
            f.c(r, 2, e["rotulo"])
            f.c(r, 3, rot_status[e["status"]], negrito=e["status"] == "faltando",
                cor="C00000" if e["status"] == "faltando" else GRAFITE)
            r += 1

    larg = {"A": 7, "B": 44, "C": 30, "D": 30, "E": 22, "F": 18, "G": 14, "H": 9, "I": 9}
    for k in range(n_p):
        larg[L(COL_PERFIL + k)] = 9
        larg[L(c_fixo + k)] = 8
        larg[L(c_unid + k)] = 8
    for k in col_f + [c_total]:
        larg[L(k)] = max(larg.get(L(k), 0), 18)
    larg[L(c_obs)] = max(larg.get(L(c_obs), 0), 30)
    larg[L(c_tot)] = 10
    larg[L(c_mo)] = 15
    larg[L(c_mo + 1)] = 3
    f.larguras(larg)
    # o bloco de produtividade fica agrupado (recolhível) para não poluir a impressão
    ws.column_dimensions.group(L(c_fixo), L(ultima), outline_level=1, hidden=False)
    ws.freeze_panes = "D4"
    _impressao(ws)
    ws.print_area = f"A1:{L(max(c_mo, c_obs))}{r}"
    q = _q(nome)
    cel = {ch: f"{q}!${T}${v}" for ch, v in pos.items()}
    cel.update({"aba": nome, "interno": interno, "obra": f"{q}!${T}${pos['preco_forn']}", "global": f"{q}!${T}${pos['preco']}",
                "fases_ext": f"{q}!{FA}", "etapas_ext": f"{q}!$F${r0}:$F${r_fim}", "horas_ext": f"{q}!{col_tot}",
                "fase": {fr["id"]: {"preco": f"{q}!${L(col_f[k])}${pos['preco']}"} for k, fr in enumerate(fases)}})
    return cel


# ----------------------------------------------------------------------------- Fornecimentos
def _aba_fornecimentos(wb, proposta, resultados, ref):
    ws = wb.create_sheet("Fornecimentos")
    f = Folha(ws)
    f.larguras({"A": 8, "B": 46, "C": 16, "D": 8, "E": 10, "F": 16, "G": 16, "H": 9, "I": 16, "J": 16, "K": 14, "L": 11, "M": 7})
    f.titulo(1, "FORNECIMENTOS E SUBCONTRATAÇÕES", 13)
    f.cabecalho(3, ["Opção", "Descrição", "Categoria", "Unid", "Qtd", "Custo unit. (R$)", "Custo (R$)", "BDI", "Preço (R$)",
                    "Fonte", "Código", "Data-base", "Fase"])
    r = 4
    linhas = {}
    for o, res in zip(proposta["opcoes"], resultados):
        for fr in res["fases"]:
            for it in fr["fornecimentos"]:
                linhas[(o["id"], it.get("uid"))] = r
                f.c(r, 1, o["id"], alinhar="center")
                f.c(r, 2, it.get("descricao", ""), quebra=True, tam=9)
                f.entrada(r, 3, it["categoria"])
                f.c(r, 4, it.get("unid", ""), alinhar="center")
                f.entrada(r, 5, float(it.get("qtd") or 0), NUM)
                f.entrada(r, 6, float(it.get("custo_unit") or 0), MOEDA)
                f.c(r, 7, f"=E{r}*F{r}", MOEDA)
                f.vinculo(r, 8, f"=INDEX({ref['bdi_valores']},MATCH(C{r},{ref['bdi_chaves']},0))", PCT)
                f.c(r, 9, f"=G{r}*(1+H{r})", MOEDA)
                f.c(r, 10, it.get("fonte", ""), tam=9)
                f.c(r, 11, it.get("codigo", ""), tam=9)
                f.c(r, 12, it.get("data_base", ""), tam=9, alinhar="center")
                f.c(r, 13, fr["id"], alinhar="center")
                r += 1
    if r == 4:
        f.nota(4, "Sem fornecimentos nesta proposta. Para incluir, use a aba 'Fornecimentos e subcontratações' do app "
                  "(categorias: material, equipamento, mo_execucao, subcontrato).", 13)
    ws.freeze_panes = "A4"
    _impressao(ws)
    return linhas


# ----------------------------------------------------------------------------- Pesquisa de preços
CRITERIOS = {"mediana": ("Mediana", "MEDIAN"), "media": ("Média", "AVERAGE"), "menor": ("Menor valor", "MIN")}


def _aba_pesquisa(wb, proposta, resultados, ref, linhas_forn):
    """Cesta de preços de cada item com referências: valores, atualização, desvio da mediana e valor adotado."""
    itens = [(o, it) for o, res in zip(proposta["opcoes"], resultados) for it in res["fornecimentos"] if it.get("referencias")]
    if not itens:
        return
    cfg = proposta["parametros"]["cesta"]
    rot_crit, func = CRITERIOS.get(cfg["criterio"], CRITERIOS["mediana"])
    ws = wb.create_sheet("Pesquisa de preços")
    f = Folha(ws)
    f.larguras({"A": 30, "B": 14, "C": 11, "D": 15, "E": 11, "F": 15, "G": 12, "H": 14, "I": 15})
    f.titulo(1, "PESQUISA DE PREÇOS (CESTA)", 9)
    f.nota(2, f"Critério: {rot_crit.lower()} das referências dentro da faixa em torno da mediana (Parametros). Recomenda-se "
              f"no mínimo {cfg['minimo_fontes']} referências por item. O fator de atualização é o do índice no momento em que a "
              "cesta foi aplicada (1 = sem atualização).", 9)
    faixa = ref["cesta_faixa"]
    r = 4
    for o, it in itens:
        f.secao(r, f"Opção {o['id']} — {it.get('descricao', '')}", 9)
        r += 1
        f.cabecalho(r, ["Fonte / fornecedor", "Código", "Data-base", "Valor (R$)", "Fator", "Atualizado (R$)",
                        "Desvio", "Situação", "Na cesta (R$)"])
        ini = r + 1
        fim = ini + len(it["referencias"]) - 1
        for k, rf_ in enumerate(it["referencias"]):
            r = ini + k
            f.c(r, 1, rf_.get("fonte", ""), tam=9)
            f.c(r, 2, rf_.get("codigo", ""), tam=9)
            f.c(r, 3, rf_.get("data_base", ""), tam=9, alinhar="center")
            f.entrada(r, 4, float(rf_.get("valor") or 0), MOEDA)
            f.entrada(r, 5, float(rf_.get("fator") or 1.0), '0.0000')
            f.c(r, 6, f"=D{r}*E{r}", MOEDA)
            f.c(r, 7, f"=IF(MEDIAN($F${ini}:$F${fim})=0,0,F{r}/MEDIAN($F${ini}:$F${fim})-1)", PCT)
            f.c(r, 8, f'=IF(ABS(G{r})>{faixa},"fora da faixa","ok")', alinhar="center")
            # MEDIAN, AVERAGE e MIN ignoram o texto vazio: só os valores dentro da faixa entram no adotado
            f.c(r, 9, f'=IF(H{r}="ok",F{r},"")', MOEDA)
        r = fim + 1
        rotulos = [("Mediana", f"=MEDIAN(F{ini}:F{fim})"), ("Média", f"=AVERAGE(F{ini}:F{fim})"), ("Menor", f"=MIN(F{ini}:F{fim})"),
                   ("Referências / dentro da faixa", f'=COUNT(F{ini}:F{fim})&" / "&COUNT(I{ini}:I{fim})')]
        for rot, form in rotulos:
            f.c(r, 8, rot, alinhar="right", tam=9)
            f.c(r, 9, form, MOEDA if rot != rotulos[-1][0] else None, alinhar="right")
            r += 1
        f.c(r, 8, f"Valor adotado ({rot_crit.lower()})", negrito=True, alinhar="right")
        adotado = f"I{r}"
        f.c(r, 9, f"=IF(COUNT(I{ini}:I{fim})=0,{func}(F{ini}:F{fim}),{func}(I{ini}:I{fim}))", MOEDA, negrito=True, fundo=FUNDO_CAB)
        r += 1
        lf = linhas_forn.get((o["id"], it.get("uid")))
        if lf:
            f.c(r, 8, "Custo unit. em Fornecimentos", alinhar="right", tam=9)
            f.vinculo(r, 9, f"=Fornecimentos!F{lf}", MOEDA)
            f.c(r + 1, 8, "Diferença para o adotado", alinhar="right", tam=9)
            f.c(r + 1, 9, f"=IF({adotado}=0,0,I{r}/{adotado}-1)", PCT)
            r += 2
        r += 1
    ws.freeze_panes = "A4"
    _impressao(ws)


# ----------------------------------------------------------------------------- Cronograma
def _aba_cronograma(wb, cat, proposta, resultados, celulas, ref):
    ws = wb.create_sheet("Cronograma")
    f = Folha(ws)
    G0 = 11                       # primeira coluna do Gantt (K)
    semanas = max([math.ceil(r["prazo_dias_uteis"] / 5) for r in resultados] + [1]) + 1
    ultima = G0 - 1 + semanas
    f.larguras({"A": 6, "B": 40, "C": 10, "D": 8, "E": 10, "F": 8, "G": 8, "H": 10, "I": 28, "J": 2})
    for k in range(semanas):
        ws.column_dimensions[L(G0 + k)].width = 3.6
    f.titulo(1, "CRONOGRAMA (DIAS ÚTEIS) E GANTT SEMANAL", min(ultima, 30))
    f.nota(2, "Dias úteis = horas da etapa ÷ (horas/dia × equipe), no mínimo 1. Fases em sequência; dentro da fase, etapas "
              "em sequência e atividades da mesma etapa em paralelo. Fases de obra (% ou pacote) duram os meses da obra "
              "(21 dias úteis ou 30 corridos por mês).", min(ultima, 30))
    r = 4
    cfg_op = {o["id"]: o for o in proposta["opcoes"]}
    for o, res in zip(proposta["opcoes"], resultados):
        cel = celulas[o["id"]]
        f.secao(r, o["nome"] + (" — estimativa interna" if res["interno"] else ""), 9)
        r += 1
        f.cabecalho(r, ["Fase", "Etapa", "Horas", "Equipe", "Dias úteis", "Início", "Fim", "Dias corridos", "Marco de conclusão"])
        for k in range(semanas):
            f.c(r, G0 + k, k + 1, INT, negrito=True, cor=AZUL, fundo=FUNDO_CAB, alinhar="center", tam=8)
        r += 1
        r_ini = r
        minimos = {fr["id"]: motor.prazos_minimos(cfg_op[o["id"]], fr) for fr in res["fases"]}
        for i, e in enumerate(res["cronograma"]):
            f.c(r, 1, e["fase"], alinhar="center")
            f.c(r, 2, e["rotulo"])
            if e.get("meses"):      # fase de obra (% ou pacote): duração em meses
                f.c(r, 3, None)
                f.c(r, 4, None)
                f.entrada(r, 5, e["dias_uteis"], INT)
                f.entrada(r, 8, e["dias_corridos"], INT)
            else:
                f.vinculo(r, 3, f"=SUMIFS({cel['horas_ext']},{cel['fases_ext']},A{r},{cel['etapas_ext']},B{r})", HORAS)
                eq = proposta["parametros"]["cronograma"].get("equipe_etapa", {}).get(e["etapa"],
                                                                                    proposta["parametros"]["cronograma"]["equipe"])
                f.entrada(r, 4, eq, INT)
                base = f"IF(C{r}>0,MAX(1,ROUNDUP(C{r}/({ref['horas_dia']}*D{r}),0)),0)"
                extra = minimos.get(e["fase"], {}).get(e["etapa"], 0)
                f.c(r, 5, f"=MAX({base},{extra})" if extra else "=" + base, INT)
                f.c(r, 8, f"=ROUNDUP(E{r}*7/5,0)", INT, alinhar="center")
            f.c(r, 6, 1 if i == 0 else f"=G{r - 1}+1", INT, alinhar="center")
            f.c(r, 7, f"=F{r}+E{r}-1", INT, alinhar="center")
            f.c(r, 9, e["marco"], tam=9)
            for k in range(semanas):
                f.c(r, G0 + k, None, borda=True)
            r += 1
        r_fim = r - 1
        if res["cronograma"]:
            faixa = f"{L(G0)}{r_ini}:{L(G0 - 1 + semanas)}{r_fim}"
            ws.conditional_formatting.add(faixa, FormulaRule(
                formula=[f"AND({L(G0)}${r_ini - 1}*5>=$F{r_ini},({L(G0)}${r_ini - 1}-1)*5+1<=$G{r_ini})"],
                fill=_fill(AZUL), stopIfTrue=False))
        f.c(r, 2, "PRAZO TOTAL", negrito=True, fundo=FUNDO_CAB)
        f.c(r, 5, f"=SUM(E{r_ini}:E{max(r_fim, r_ini)})", INT, negrito=True, fundo=FUNDO_CAB)
        f.c(r, 8, f"=SUM(H{r_ini}:H{max(r_fim, r_ini)})", INT, negrito=True, fundo=FUNDO_CAB)
        cel["prazo_uteis"] = f"Cronograma!$E${r}"
        cel["prazo_corridos"] = f"Cronograma!$H${r}"
        r += 2
    ws.freeze_panes = "C4"
    _impressao(ws)


# ----------------------------------------------------------------------------- Pagamento
def _aba_pagamento(wb, proposta, resultados, celulas):
    ws = wb.create_sheet("Pagamento")
    f = Folha(ws)
    f.larguras({"A": 76, "B": 10, "C": 18})
    f.titulo(1, "FORMA DE PAGAMENTO POR FASE", 3)
    r = 3
    for o, res in zip(proposta["opcoes"], resultados):
        if res["interno"]:
            continue
        cel = celulas[o["id"]]
        f.secao(r, o["nome"], 3)
        r += 1
        for fr in res["fases"]:
            f.cabecalho(r, [f"{fr['nome']} ({motor.ROTULO_COBRANCA[fr['cobranca']]}) — marco", "%", "Valor (R$)"])
            r += 1
            r_ini = r
            for m in fr["pagamento"]:
                f.c(r, 1, m["marco"], quebra=True)
                f.entrada(r, 2, float(m["pct"]), PCT)
                f.c(r, 3, f"=B{r}*{cel['fase'][fr['id']]['preco']}", MOEDA)
                r += 1
            f.c(r, 1, "TOTAL DA FASE (deve somar 100%)", negrito=True, fundo=FUNDO_CAB)
            f.c(r, 2, f"=SUM(B{r_ini}:B{max(r - 1, r_ini)})", PCT, negrito=True, fundo=FUNDO_CAB)
            f.c(r, 3, f"=SUM(C{r_ini}:C{max(r - 1, r_ini)})", MOEDA, negrito=True, fundo=FUNDO_CAB)
            ws.conditional_formatting.add(f"B{r}", FormulaRule(formula=[f"ABS(B{r}-1)>0.0001"], fill=_fill("F8CBAD")))
            r += 2
    f.nota(r, "Nas fases por % sobre as medições, o valor é estimado sobre o valor global da obra; o faturado acompanha "
              "as medições aprovadas.", 3)
    _impressao(ws, paisagem=False)


# ----------------------------------------------------------------------------- Resumo
def _aba_resumo(ws, cat, proposta, resultados, celulas):
    f = Folha(ws)
    n = len(resultados)
    ate = 2 + n
    f.larguras({"A": 46, "B": 4, **{L(3 + i): 24 for i in range(n)}})
    _logo(ws, "A1")
    for linha in range(1, 5):
        ws.row_dimensions[linha].height = 18
    org = _org()
    f.titulo(6, f"PROPOSTA COMERCIAL {proposta['numero']} — RESUMO DAS OPÇÕES", ate)
    cl = proposta["cliente"]
    info = [
        ("Cliente", cl.get("razao", "")),
        ("Objeto", proposta.get("objeto", "")),
        ("Endereço do objeto", cl.get("endereco", "")),
        ("Data de emissão / validade", f"{_data_br(proposta['data'])} — {proposta['validade_dias']} dias corridos"),
        ("Responsável técnico", f"{org.get('resp', '').title()} — {org.get('registros', '')}"),
        ("Empresa", f"{org.get('fantasia', '')} — CNPJ {org.get('cnpj', '')}"),
    ]
    r = 7
    for rot, val in info:
        f.c(r, 1, rot, negrito=True, borda=False)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=ate)
        f.c(r, 3, val, borda=False, quebra=True)
        r += 1
    r += 1
    f.cabecalho(r, ["Indicador", ""] + [o["nome"] for o in proposta["opcoes"]])
    r += 1
    textos = [
        ("Modalidade", lambda res: "Estimativa interna (não ofertada)" if res["interno"] else "Proposta em fases"),
        ("Fases e cobranças", lambda res: " · ".join(f"{x['nome']}: {motor.ROTULO_COBRANCA[x['cobranca']]}"
                                                     for x in res["fases"]) if not res["interno"] else "—"),
    ]
    for rot, fn in textos:
        f.c(r, 1, rot)
        f.c(r, 2, None)
        for i, res in enumerate(resultados):
            f.c(r, 3 + i, fn(res), alinhar="right", quebra=True, tam=9)
        r += 1
    linhas = [
        ("Horas técnicas", "h_fase", HORAS),
        ("Prazo (dias úteis)", "prazo_uteis", INT),
        ("Prazo (dias corridos)", "prazo_corridos", INT),
        ("Custo dos serviços", "custo_fase", MOEDA),
        ("Preço — parcela fixa", "p_fixo", MOEDA),
        ("Preço — parcela estimada", "p_est", MOEDA),
        ("PREÇO TOTAL (FIXO + ESTIMADO)", "preco", MOEDA),
        ("Taxa de desempenho — potencial (fora do preço)", "de_pot", MOEDA),
        ("Referência: tudo em escopo fechado", "fechado", MOEDA),
        ("Preço mínimo dos serviços (margem zero)", "minimo", MOEDA),
        ("Margem efetiva", "me", PCT),
        ("Preço dos serviços por hora (R$/h)", "rs_h", MOEDA),
        ("Preço por m² (R$/m²)", "rs_m2", MOEDA),
        ("Referência: % do CUB × área", "ref_cub", MOEDA),
    ]
    for rot, chave, fmt in linhas:
        forte = rot.isupper()
        f.c(r, 1, rot, negrito=forte, fundo=FUNDO_CAB if forte else None)
        f.c(r, 2, None, fundo=FUNDO_CAB if forte else None)
        for i, o in enumerate(proposta["opcoes"]):
            ref_cel = celulas[o["id"]].get(chave)
            if ref_cel:
                f.vinculo(r, 3 + i, "=" + ref_cel, fmt, negrito=forte)
            else:
                f.c(r, 3 + i, "—")
            ws.cell(r, 3 + i).alignment = Alignment(horizontal="right")
        r += 1
    r += 1
    f.secao(r, "Preço por fase", ate)
    r += 1
    f.cabecalho(r, ["Opção / fase", "", "Cobrança", "Preço da fase"])
    r += 1
    for o, res in zip(proposta["opcoes"], resultados):
        if res["interno"]:
            continue
        for fr in res["fases"]:
            f.c(r, 1, f"{o['nome']} — {fr['nome']}")
            f.c(r, 2, None)
            f.c(r, 3, motor.ROTULO_COBRANCA[fr["cobranca"]])
            f.vinculo(r, 4, "=" + celulas[o["id"]]["fase"][fr["id"]]["preco"], MOEDA)
            r += 1
    r += 1
    f.nota(r, cat.get("nota", ""), ate)
    f.nota(r + 1, "Valores com tributos incidentes. Parcela estimada: hora técnica e % sobre as medições da obra. Fundo "
                  "amarelo = entrada editável; verde = vínculo com outra aba. O documento técnico só produz efeito "
                  "acompanhado da ART registrada e quitada (Lei 6.496/1977).", ate)
    ws.freeze_panes = "A7"
    _impressao(ws, paisagem=n > 2)
    ws.page_setup.fitToHeight = 1


def _data_br(iso):
    try:
        a, m, d = iso.split("-")
        return f"{d}/{m}/{a}"
    except (ValueError, AttributeError):
        return str(iso)


# ----------------------------------------------------------------------------- API
def gerar(cat, proposta, resultado):
    """Bytes do .xlsx com todas as opções da proposta."""
    resultados = resultado["opcoes"]
    wb = Workbook()
    resumo = wb.active
    resumo.title = "Resumo"
    ref = _aba_parametros(wb, cat, proposta)
    celulas = {}
    # estimativas primeiro: a base do % das fases de obra aponta para as células delas
    pares = list(zip(proposta["opcoes"], resultados))
    for o, res in sorted(pares, key=lambda par: not par[1]["interno"]):
        celulas[o["id"]] = _aba_opcao(wb, cat, proposta, o, res, ref, celulas)
    linhas_forn = _aba_fornecimentos(wb, proposta, resultados, ref)
    _aba_pesquisa(wb, proposta, resultados, ref, linhas_forn)
    _aba_cronograma(wb, cat, proposta, resultados, celulas, ref)
    _aba_pagamento(wb, proposta, resultados, celulas)
    _aba_resumo(resumo, cat, proposta, resultados, celulas)
    ordem = ["Resumo"] + [celulas[o["id"]]["aba"] for o in proposta["opcoes"]]
    wb._sheets.sort(key=lambda ws: (ordem.index(ws.title) if ws.title in ordem else len(ordem)))
    wb.move_sheet("Parametros", offset=len(wb.sheetnames))
    wb.active = 0
    wb.properties.title = f"Proposta {proposta['numero']}"
    wb.properties.creator = _org().get("fantasia", "KAPTEK")
    wb.calculation.fullCalcOnLoad = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
