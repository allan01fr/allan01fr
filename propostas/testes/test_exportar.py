import io
import shutil
import subprocess

import pytest
from docx import Document
from openpyxl import load_workbook

from conftest import EXEMPLOS, abrir_exemplo
from orcamento import exportar_docx, exportar_xlsx, motor

ROTULOS = {"Horas técnicas": "total_h", "Prazo (dias corridos)": "prazo_dias_corridos",
           "Preço — parcela fixa": "preco_fixo", "Preço — parcela estimada": "preco_estimado",
           "PREÇO TOTAL (FIXO + ESTIMADO)": "preco", "Referência: tudo em escopo fechado": "preco_fechado",
           "Preço mínimo dos serviços (margem zero)": "preco_minimo_servicos"}


@pytest.mark.parametrize("caminho", EXEMPLOS)
def test_xlsx_tem_formulas_e_abas(caminho, cat, params):
    p = abrir_exemplo(caminho, cat, params)
    res = motor.precificar(cat, p)
    wb = load_workbook(io.BytesIO(exportar_xlsx.gerar(cat, p, res)))
    esperadas = {"Resumo", "Parametros", "Fornecimentos", "Cronograma", "Pagamento"}
    esperadas |= {f"Opção {o['id']}" for o in p["opcoes"]}
    assert esperadas <= set(wb.sheetnames)
    formulas = sum(1 for ws in wb for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("="))
    assert formulas > 50


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice não instalado")
@pytest.mark.parametrize("caminho", EXEMPLOS)
def test_xlsx_recalculado_bate_com_o_motor(caminho, cat, params, tmp_path):
    p = abrir_exemplo(caminho, cat, params)
    res = motor.precificar(cat, p)
    arq = tmp_path / "proposta.xlsx"
    arq.write_bytes(exportar_xlsx.gerar(cat, p, res))
    saida = tmp_path / "rec"
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(saida), str(arq)],
                   check=True, capture_output=True, timeout=180)
    wb = load_workbook(saida / "proposta.xlsx", data_only=True)
    codigos = ("#NAME?", "#VALUE!", "#REF!", "#DIV/0!", "#N/A", "#NUM!", "#NULL!", "Err:")
    erros = [f"{ws.title}!{c.coordinate}" for ws in wb for row in ws.iter_rows() for c in row
             if isinstance(c.value, str) and c.value.startswith(codigos)]
    assert not erros
    ws = wb["Resumo"]
    linha = {ws.cell(r, 1).value: r for r in range(1, ws.max_row + 1)}
    for i, r in enumerate(res["opcoes"]):
        for rot, chave in ROTULOS.items():
            assert float(ws.cell(linha[rot], 3 + i).value or 0) == pytest.approx(r[chave], abs=0.01), (r["nome"], rot)
        # matriz por fase
        wo = wb[f"Opção {r['id']}"]
        rot_o = {wo.cell(k, 2).value: k for k in range(1, wo.max_row + 1) if wo.cell(k, 2).value}
        for k, fr in enumerate(r["fases"]):
            for rot, chave in (("PREÇO DA FASE", "preco"), ("Custo da fase", "custo_servicos"), ("Horas da fase", "total_h")):
                assert float(wo.cell(rot_o[rot], 3 + k).value or 0) == pytest.approx(fr[chave], abs=0.01), (r["id"], fr["id"], rot)
    # cesta de preços: o valor adotado na planilha é o custo unitário aplicado ao item
    com_cesta = [f for o in p["opcoes"] for f in o["fornecimentos"] if f.get("referencias")]
    if com_cesta:
        pq = wb["Pesquisa de preços"]
        adotados = [pq.cell(r, 9).value for r in range(1, pq.max_row + 1) if str(pq.cell(r, 8).value).startswith("Valor adotado")]
        diferencas = [pq.cell(r, 9).value for r in range(1, pq.max_row + 1) if pq.cell(r, 8).value == "Diferença para o adotado"]
        assert adotados == pytest.approx([f["custo_unit"] for f in com_cesta], abs=0.01)
        assert diferencas == pytest.approx([0.0] * len(com_cesta), abs=1e-6)


@pytest.mark.parametrize("caminho", EXEMPLOS)
def test_docx_preenche_o_template(caminho, cat, params):
    p = abrir_exemplo(caminho, cat, params)
    res = motor.precificar(cat, p)
    docs = exportar_docx.gerar(cat, p, res)
    ofertadas = [r for r in res["opcoes"] if not r["interno"]]       # a estimativa interna não vira proposta
    assert len(docs) == len(ofertadas)
    for (conteudo, nome, pend), r in zip(docs, ofertadas):
        assert nome.startswith("KAPTEK-proposta-") and nome.endswith(".docx")
        assert exportar_docx.verificar(conteudo) == {"margem_sup_cm": 3.2, "margem_esq_cm": 3.0, "logo": True}
        assert set(pend) <= {"[Nome e cargo do signatário]"}
        d = Document(io.BytesIO(conteudo))
        escopo, prazos, honor = d.tables[2], d.tables[3], d.tables[4]
        servicos = len({l["item"] for l in r["linhas"]})
        assert len(escopo.rows) - 1 == servicos
        assert len(prazos.rows) - 1 == (len(r["cronograma"]) if len(r["fases"]) == 1 else len(r["fases"])) + 1
        assert honor.rows[-1].cells[-1].text == exportar_docx._total_texto(r)
        texto = "\n".join(par.text for par in d.paragraphs)
        assert p["numero"] in "\n".join(c.text for row in d.tables[0].rows for c in row.cells)
        assert "[" not in texto.replace("[Nome e cargo do signatário]", "")


def _todas_as_cobrancas(cat, params):
    """Fechado com fornecimento, hora técnica, % (parte fixa + medições) e pacote total sem serviços com taxa de desempenho."""
    import copy

    from orcamento import catalogo as ct
    from orcamento import proposta as pp
    p = pp.completar(pp.padrao(copy.deepcopy(params)), cat, params)
    o = p["opcoes"][0]
    o["fases"] = [pp.nova_fase("F1", params, "Laudo"), pp.nova_fase("F2", params, "Simulação", "hora"),
                  pp.nova_fase("F3", params, "Assessoria", "percentual"), pp.nova_fase("F4", params, "Fiscalização", "pacote")]
    o["itens"] = [pp.novo_item(ct.servico(cat, "hvac.asbuilt"), "F1"), pp.novo_item(ct.servico(cat, "hvac.cfd"), "F2"),
                  pp.novo_item(ct.servico(cat, "geral.planejamento_obra"), "F3")]
    o["fornecimentos"] = [pp.novo_fornecimento("Registradores (locação)", "material", "mês", 2, 1800.0, fase="F1")]
    f3, f4 = o["fases"][2], o["fases"][3]
    f3["percentual"].update(pct=0.035, forma="fixo_mais_medicoes", parcela_fixa=3000.0)
    f3["obra"].update(base_valor=500000.0, meses=3, curva="linear")
    f4["pacote"].update(forma="total", horas=120.0, validade_meses=6, parcelas=3, desconto=0.05, excedente_pct=0.1)
    f4["obra"].update(base_valor=800000.0)
    f4["desempenho"].update(ativo=True, cenario={"economia": 0.04, "dias_antecipados": 10, "retrabalho_ok": False})
    return p


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice não instalado")
def test_xlsx_recalculado_com_todas_as_cobrancas(cat, params, tmp_path):
    p = _todas_as_cobrancas(cat, params)
    res = motor.precificar(cat, p)
    arq = tmp_path / "fases.xlsx"
    arq.write_bytes(exportar_xlsx.gerar(cat, p, res))
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(tmp_path / "rec"), str(arq)],
                   check=True, capture_output=True, timeout=180)
    wb = load_workbook(tmp_path / "rec" / "fases.xlsx", data_only=True)
    codigos = ("#NAME?", "#VALUE!", "#REF!", "#DIV/0!", "#N/A", "#NUM!", "#NULL!", "Err:")
    assert not [c.coordinate for ws in wb for row in ws.iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith(codigos)]
    r = res["opcoes"][0]
    ws = wb["Opção A"]
    rot = {ws.cell(k, 2).value: k for k in range(1, ws.max_row + 1) if ws.cell(k, 2).value}
    campos = (("PREÇO DA FASE", "preco"), ("Preço — parcela fixa", "preco_fixo"), ("Preço — parcela estimada", "preco_estimado"),
              ("Custo da fase", "custo_servicos"), ("Horas da fase", "total_h"), ("Margem efetiva", "margem_efetiva"))
    for k, fr in enumerate(r["fases"]):
        for rotulo, chave in campos:
            assert float(ws.cell(rot[rotulo], 3 + k).value or 0) == pytest.approx(fr[chave], abs=0.01), (fr["id"], rotulo)
    pot = ws.cell(rot["TAXA DE DESEMPENHO — POTENCIAL"], 6).value
    assert pot == pytest.approx(r["fases"][3]["desempenho"]["potencial"], abs=0.01)
    wr = wb["Resumo"]
    linha = {wr.cell(k, 1).value: k for k in range(1, wr.max_row + 1)}
    for rotulo, chave in ROTULOS.items():
        assert float(wr.cell(linha[rotulo], 3).value or 0) == pytest.approx(r[chave], abs=0.01), rotulo
