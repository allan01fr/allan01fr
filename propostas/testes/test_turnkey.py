import copy
import io
import os
import re

import pytest
from docx import Document

from conftest import RAIZ, abrir_exemplo
from orcamento import catalogo as ct
from orcamento import exportar_docx, motor
from orcamento import proposta as pp

EXEMPLO = os.path.join(RAIZ, "exemplos", "hvac_turnkey.json")


@pytest.fixture
def p(cat, params):
    return abrir_exemplo(EXEMPLO, cat, params)


def opcao_a(cat, p):
    return motor.precificar(cat, p)["opcoes"][0]


def test_exemplo_separa_modalidade_de_sistema(p):
    a, b = p["opcoes"]
    assert a["modalidade"] == "turnkey" and b["modalidade"] == "servicos"
    projeto = next(it for it in a["itens"] if it["servico"] == "hvac.projeto")
    assert projeto["dados"]["sistema"] == "vrf" and projeto["dados"]["ventilacao_tipo"] == "doas"
    assert "turnkey" not in p["objeto"].lower()       # turnkey é a modalidade da opção, não o objeto


def test_turnkey_sempre_preco_global(cat, p):
    p["opcoes"][0]["modelo_preco"] = "hora"
    r = opcao_a(cat, p)
    assert r["modelo_preco"] == "fechado" and r["preco"] == pytest.approx(r["preco_fechado"])


def test_provisao_de_garantia(cat, p):
    r = opcao_a(cat, p)
    g = next(d for d in r["diretos"] if d.get("garantia"))
    assert g["valor"] == pytest.approx(r["custo_fornec"] * p["parametros"]["turnkey"]["garantia_pct"])
    p["opcoes"][0]["modalidade"] = "servicos"
    assert not any(d.get("garantia") for d in opcao_a(cat, p)["diretos"])


def test_checklist_e_avisos(cat, params):
    p = pp.completar(pp.padrao(copy.deepcopy(params)), cat, params)
    o = p["opcoes"][0]
    o["modalidade"] = "turnkey"
    o["itens"] = [pp.novo_item(ct.servico(cat, "hvac.projeto"))]
    res = motor.precificar(cat, p)
    estado = {c["id"]: c["status"] for c in res["opcoes"][0]["checklist_turnkey"]}
    assert estado["projeto"] == "atendida" and estado["aprovacoes"] == "opcional"
    assert estado["equipamentos"] == estado["execucao"] == estado["comissionamento"] == "faltando"
    assert any("Gestão da obra" in a for a in res["avisos"])
    o["turnkey_dispensas"] = ["gestao"]
    res = motor.precificar(cat, p)
    assert {c["id"]: c["status"] for c in res["opcoes"][0]["checklist_turnkey"]}["gestao"] == "dispensada"
    assert not any("Gestão da obra" in a for a in res["avisos"])


def test_exemplo_completo_sem_avisos(cat, p):
    res = motor.precificar(cat, p)
    assert all(c["status"] in ("atendida", "opcional") for c in res["opcoes"][0]["checklist_turnkey"])
    assert res["avisos"] == []


def test_etapa_de_suprimentos_com_lead_time(cat, p):
    a, b = motor.precificar(cat, p)["opcoes"]
    etapas = {e["etapa"]: e for e in a["cronograma"]}
    assert etapas["suprimentos"]["dias_uteis"] == p["opcoes"][0]["prazo_fornecimento_dias"]
    assert etapas["suprimentos"]["fim"] < etapas["obra"]["inicio"]
    assert "suprimentos" not in {e["etapa"] for e in b["cronograma"]}


def test_pagamento_turnkey(cat, p):
    r = opcao_a(cat, p)
    pcts = [m["pct"] for m in r["pagamento"]]
    assert sum(pcts) == pytest.approx(1.0)
    assert pcts[0] == p["parametros"]["pagamento"]["sinal"]
    assert "equipamentos" in r["pagamento"][1]["marco"]
    assert pcts[-1] >= p["parametros"]["turnkey"]["retencao_aceite"] - 1e-9
    assert all(round(x * 100) % 5 == 0 for x in pcts)
    p["parametros"]["turnkey"]["retencao_aceite"] = 0.25
    assert opcao_a(cat, p)["pagamento"][-1]["pct"] >= 0.25 - 1e-9


def _numero(texto):
    return float(re.sub(r"[^\d,]", "", texto).replace(",", "."))


def test_docx_turnkey_com_eventos_por_etapa(cat, p):
    res = motor.precificar(cat, p)
    conteudo, _, _ = exportar_docx.gerar(cat, p, res)[0]
    doc = Document(io.BytesIO(conteudo))
    texto = "\n".join(par.text for par in doc.paragraphs)
    assert "turnkey" in texto.lower()
    tab = doc.tables[4]
    linhas = [[c.text for c in row.cells] for row in tab.rows]
    descricoes = [l[0] for l in linhas]
    assert "Fornecimento de equipamentos" in descricoes and "Execução e instalação" in descricoes
    assert not any("evaporadoras" in d for d in descricoes)            # itens e unitários ficam só na planilha
    assert any(l[0] == "Taxa de ART junto ao CREA" and l[-1] == "inclusa" for l in linhas)
    eventos = [l for l in linhas[1:] if l[1] == "1 vb"]
    assert sum(_numero(l[-1]) for l in eventos) == pytest.approx(res["opcoes"][0]["preco_fechado"], abs=0.01)
    assert f"{p['parametros']['turnkey']['garantia_meses']} meses" in texto
