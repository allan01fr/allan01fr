import io
import math
import os
import shutil
import subprocess

import pytest
from openpyxl import load_workbook

from conftest import RAIZ, abrir_exemplo
from orcamento import exportar_viabilidade, motor
from orcamento import viabilidade as vb

C, A, I, N = 200000.0, 30000.0, 0.10, 15      # CAPEX, economia anual, TMA, horizonte


def estudo(capex=C, reducao=0.30, horizonte=N, ge=0.0, gm=0.0, manut=None, **base):
    """Linha de base de 100 MWh/ano a R$ 1/kWh: redução de 30% = R$ 30 mil/ano."""
    v = vb.padrao()
    v["base"].update({"modo": "kwh", "kwh_ano": 100000.0, "tarifa": 1.0, "manutencao_ano": 50000.0, "paliativos": 0.0, **base})
    v["financeiro"].update({"tma": I, "horizonte": horizonte, "reajuste_energia": ge, "reajuste_manutencao": gm})
    v["alternativas"] = [{**vb.nova_alternativa("X"), "capex": capex, "reducao_pct": reducao, "manutencao_ano": manut}]
    return v


def alt(v, precos=None):
    return vb.calcular(v, precos)["alternativas"][0]


# ----------------------------------------------------------------------------- fórmulas fechadas (sem reajuste)
def test_breakeven_e_vpl_de_anuidade():
    a = alt(estudo())
    assert a["economia_ano1"] == pytest.approx(A)
    assert a["breakeven"] == pytest.approx(C / A)              # as contas empatam: economia acumulada = CAPEX
    assert a["breakeven_meses"] == math.ceil(C / A * 12)
    assert a["vpl"] == pytest.approx(A * (1 - (1 + I) ** -N) / I - C)
    assert a["classificacao"] == "Viável"


def test_payback_e_o_retorno_descontado_pela_tma():
    a = alt(estudo())
    fechado = -math.log(1 - C * I / A) / math.log(1 + I)
    # a interpolação linear dentro do ano difere da forma contínua em centésimos de ano
    assert a["payback"] == pytest.approx(fechado, abs=0.05)    # o projeto se paga, remunerando o capital
    assert a["payback_meses"] == math.ceil(a["payback"] * 12)
    assert a["payback"] > a["breakeven"]


def test_tir_zera_o_vpl():
    a = alt(estudo())
    assert vb.vpl(a["fluxo"], a["tir"]) == pytest.approx(0, abs=1e-4)
    assert a["tir"] > I                       # VPL positivo ⇒ TIR acima da TMA


def test_economia_minima_zera_o_vpl_e_margem():
    a = alt(estudo())
    ajustada = alt(estudo(reducao=0.30 * a["economia_minima"] / A))
    assert ajustada["vpl"] == pytest.approx(0, abs=1e-6)
    assert a["margem_seguranca"] == pytest.approx(a["vp_economia"] / C - 1)


def test_sensibilidade_exata():
    v = estudo(manut=40000.0)                 # economia de energia 30 mil + manutenção 10 mil
    a = alt(v)
    v_tarifa = estudo(manut=40000.0, tarifa=1.1)
    assert a["sensibilidade"]["tarifa"][0.1] == pytest.approx(alt(v_tarifa)["vpl"])
    assert a["sensibilidade"]["capex"][-0.2] == pytest.approx(alt(estudo(capex=0.8 * C, manut=40000.0))["vpl"])
    assert a["sensibilidade"]["economia"][0.0] == pytest.approx(a["vpl"])


def test_reajuste_entra_no_fluxo():
    a = alt(estudo(ge=0.05, gm=0.03, manut=40000.0))
    assert a["fluxo"][1] == pytest.approx(30000 + 10000)
    assert a["fluxo"][3] == pytest.approx(30000 * 1.05 ** 2 + 10000 * 1.03 ** 2)


# ----------------------------------------------------------------------------- horizonte e casos-limite
def test_corte_no_horizonte():
    a = alt(estudo(horizonte=5))              # 5 × 30 mil < 200 mil
    assert len(a["fluxo"]) == 6
    assert a["breakeven"] is None and a["breakeven_meses"] is None
    assert a["payback"] is None and a["payback_meses"] is None
    assert a["classificacao"] == "Inviável financeiramente"
    assert len(alt(estudo(horizonte=80))["fluxo"]) == vb.HORIZONTE_MAX + 1


def test_sem_economia_e_quick_win():
    s = alt(estudo(reducao=0.0))
    assert s["classificacao"] == "Sem economia" and s["breakeven"] is None and s["tir"] is None
    q = alt(estudo(capex=45000.0))
    assert q["breakeven"] == pytest.approx(1.5) and q["classificacao"] == "Quick win"
    assert alt(estudo(capex=0.0))["breakeven"] == 0.0 and alt(estudo(capex=0.0))["payback"] == 0.0


def test_geracao_fv_nao_deixa_consumo_negativo():
    v = estudo(reducao=0.5)
    v["alternativas"][0]["geracao_kwh"] = 80000.0
    a = alt(v)
    assert a["consumo"] == 0.0 and a["economia_energia_ano1"] == pytest.approx(100000.0)


def test_linha_de_base_por_tr_e_perda_de_desempenho():
    v = vb.padrao()
    v["base"].update({"modo": "tr", "tr": 75, "kw_tr_real": 1.45, "kw_tr_projeto": 1.0, "horas_ano": 3750, "fator_carga": 0.65,
                      "tarifa": 0.85, "manutencao_ano": 60000, "paliativos": 25000})
    e = vb.calcular(v)
    assert e["consumo_base"] == pytest.approx(75 * 1.45 * 3750 * 0.65)
    assert e["excesso_vs_projeto"] == pytest.approx(0.45)
    assert e["perda_desempenho"] == pytest.approx(75 * 0.45 * 3750 * 0.65 * 0.85)
    fd = [(1 + v["financeiro"]["reajuste_energia"]) ** (t - 1) / (1 + v["financeiro"]["tma"]) ** t for t in range(1, 16)]
    fm = [(1 + v["financeiro"]["reajuste_manutencao"]) ** (t - 1) / (1 + v["financeiro"]["tma"]) ** t for t in range(1, 16)]
    assert e["lcc_base"] == pytest.approx(25000 + e["custo_energia_base"] * sum(fd) + 60000 * sum(fm))


def test_capex_vinculado_ao_preco_da_opcao():
    v = estudo()
    v["alternativas"][0].update({"capex_opcao": "A", "capex_extra": 5000.0})
    a = alt(v, {"A": 150000.0})
    assert a["capex"] == pytest.approx(155000.0) and a["capex_origem"] == "preço da opção A + extra"
    assert alt(v, {"B": 1.0})["capex"] == pytest.approx(C)      # opção inexistente: vale o CAPEX informado


def test_exemplo_turnkey_usa_o_preco_da_opcao(cat, params):
    p = abrir_exemplo(os.path.join(RAIZ, "exemplos", "hvac_turnkey.json"), cat, params)
    res = motor.precificar(cat, p)
    e = vb.calcular(p["viabilidade"], {r["id"]: r["preco"] for r in res["opcoes"]})
    turnkey = next(a for a in e["alternativas"] if "opção A" in a["capex_origem"])
    assert turnkey["capex"] == pytest.approx(res["opcoes"][0]["preco_fechado"])


# ----------------------------------------------------------------------------- planilha
def _recalcular(conteudo, tmp_path, nome):
    arq = tmp_path / nome
    arq.write_bytes(conteudo)
    saida = tmp_path / "rec"
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(saida), str(arq)],
                   check=True, capture_output=True, timeout=180)
    return load_workbook(saida / nome, data_only=True)


def _sem_erros(wb):
    codigos = ("#NAME?", "#VALUE!", "#REF!", "#DIV/0!", "#N/A", "#NUM!", "#NULL!", "Err:")
    erros = [f"{ws.title}!{c.coordinate}={c.value}" for ws in wb for row in ws.iter_rows() for c in row
             if isinstance(c.value, str) and c.value.startswith(codigos)]
    assert not erros, erros[:10]


def test_planilha_tem_formulas_e_abas(cat, params):
    p = abrir_exemplo(os.path.join(RAIZ, "exemplos", "hvac_turnkey.json"), cat, params)
    e = vb.calcular(p["viabilidade"], {r["id"]: r["preco"] for r in motor.precificar(cat, p)["opcoes"]})
    wb = load_workbook(io.BytesIO(exportar_viabilidade.gerar(p, e)))
    assert {"Resumo", "Premissas", "Alternativas", "FluxoCaixa", "Sensibilidade"} <= set(wb.sheetnames)
    assert wb["FluxoCaixa"]._charts


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice não instalado")
def test_planilha_recalculada_bate_com_o_motor(cat, params, tmp_path):
    p = abrir_exemplo(os.path.join(RAIZ, "exemplos", "hvac_turnkey.json"), cat, params)
    v = p["viabilidade"]
    # casos-limite junto com o exemplo: sem economia e retorno além do horizonte
    v["alternativas"] += [{**vb.nova_alternativa("Sem economia"), "capex": 10000.0},
                          {**vb.nova_alternativa("Lenta"), "capex": 5e6, "reducao_pct": 0.05}]
    e = vb.calcular(v, {r["id"]: r["preco"] for r in motor.precificar(cat, p)["opcoes"]})
    wb = _recalcular(exportar_viabilidade.gerar(p, e), tmp_path, "viabilidade.xlsx")
    _sem_erros(wb)
    ws = wb["Resumo"]
    linha = {ws.cell(r, 1).value: r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value}
    campos = {"CAPEX": "capex", "Economia no ano 1": "economia_ano1", "VPL": "vpl", "TIR (a.a.)": "tir",
              "Breakeven — as contas empatam": "breakeven", "Breakeven (meses)": "breakeven_meses",
              "Payback — o projeto se paga (descontado pela TMA)": "payback", "Payback (meses)": "payback_meses",
              "LCC no horizonte (paliativos + CAPEX + OPEX a VP)": "lcc",
              "CLASSIFICAÇÃO": "classificacao"}
    for k, a in enumerate(e["alternativas"]):
        for rot, chave in campos.items():
            cel = ws.cell(linha[rot], 3 + k).value
            esperado = a[chave]
            if chave == "classificacao":
                assert cel == esperado, (a["nome"], rot)
            elif esperado is None:
                assert cel in ("—", "> horizonte"), (a["nome"], rot, cel)
            else:
                assert cel == pytest.approx(esperado, rel=1e-6, abs=1e-6), (a["nome"], rot)
