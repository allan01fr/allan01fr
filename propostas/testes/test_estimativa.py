import io
import os

import pytest
from docx import Document

from conftest import RAIZ, abrir_exemplo
from orcamento import exportar_docx, motor
from orcamento import viabilidade as vb

EXEMPLO = os.path.join(RAIZ, "exemplos", "hvac_retrofit_fases.json")


@pytest.fixture
def p(cat, params):
    return abrir_exemplo(EXEMPLO, cat, params)


@pytest.fixture
def res(cat, p):
    return motor.precificar(cat, p)


def por_id(res):
    return {r["id"]: r for r in res["opcoes"]}


def test_exemplo_separa_estimativa_das_fases(p):
    op = {o["id"]: o for o in p["opcoes"]}
    assert op["E"]["modalidade"] == "estimativa"
    assert op["A"]["modalidade"] == op["B"]["modalidade"] == "fases"
    assert [f["cobranca"] for f in op["A"]["fases"]] == ["fechado", "fechado", "fechado", "percentual"]
    assert op["B"]["fases"][3]["cobranca"] == "pacote"
    projeto = next(it for it in op["A"]["itens"] if it["servico"] == "hvac.projeto")
    assert projeto["dados"]["sistema"] == "vrf" and projeto["dados"]["ventilacao_tipo"] == "doas"


def test_estimativa_tem_garantia_suprimentos_e_checklist(p, res):
    e = por_id(res)["E"]
    assert e["interno"] and e["fases"][0]["pagamento"] == []
    g = next(d for d in e["diretos"] if d.get("garantia"))
    assert g["valor"] == pytest.approx(e["custo_fornec"] * p["parametros"]["turnkey"]["garantia_pct"])
    etapas = {x["etapa"]: x for x in e["cronograma"]}
    assert etapas["suprimentos"]["dias_uteis"] == 60 and etapas["suprimentos"]["fim"] < etapas["obra"]["inicio"]
    estado = {c["id"]: c["status"] for c in e["checklist_turnkey"]}
    assert estado["projeto"] == estado["gestao"] == "dispensada"
    assert "faltando" not in estado.values()
    assert res["avisos"] == []


def test_base_do_percentual_e_capex_vem_da_estimativa(p, res):
    r = por_id(res)
    f4 = r["A"]["fases"][3]
    assert f4["base"] == pytest.approx(r["E"]["valor_obra"])
    assert f4["preco_estimado"] == pytest.approx(0.04 * r["E"]["valor_obra"])
    estudo = vb.calcular(p["viabilidade"], {x["id"]: x["preco"] for x in res["opcoes"]})
    alt = next(a for a in estudo["alternativas"] if a["nome"].startswith("Retrofit VRF"))
    assert alt["capex"] == pytest.approx(r["E"]["preco"])


def test_estimativa_nao_gera_proposta(cat, p, res):
    nomes = [n for _, n, _ in exportar_docx.gerar(cat, p, res)]
    assert len(nomes) == 2 and all(("-opA-" in n) or ("-opB-" in n) for n in nomes)


def _doc(cat, p, res, opcao):
    for conteudo, nome, _ in exportar_docx.gerar(cat, p, res):
        if f"-op{opcao}-" in nome:
            return Document(io.BytesIO(conteudo))
    raise AssertionError(opcao)


def _listas(doc):
    return [x.text for x in doc.paragraphs if x.text.startswith("—")]


def test_docx_em_fases_com_percentual_e_taxa_de_desempenho(cat, p, res):
    doc = _doc(cat, p, res, "A")
    a = por_id(res)["A"]
    linhas = [[c.text for c in row.cells] for row in doc.tables[4].rows]
    descr = [l[0] for l in linhas]
    assert sum(d.startswith("Fase ") for d in descr) == 5          # 4 fases + taxa de desempenho
    assert any("4% sobre o valor de cada medição" in d for d in descr)
    assert any(l[2].startswith("até R$") and l[3] == "variável" for l in linhas)
    assert "Taxa de ART junto ao CREA (uma por fase)" in descr
    assert linhas[-1][-1] == f"{exportar_docx.moeda(a['preco_fixo'])} + {exportar_docx.moeda(a['preco_estimado'])} estimados"
    escopo = [row.cells[0].text for row in doc.tables[2].rows][1:]
    assert escopo[0] == "1.1" and escopo[-1] == "4.2"
    prazos = [row.cells[1].text for row in doc.tables[3].rows]
    assert "durante a obra (4 meses)" in prazos
    textos = " ".join(_listas(doc))
    assert "As fases são sequenciais" in textos
    assert "Fase 4 (Assessoria de obra): mensalmente, 4% sobre o valor de cada medição aprovada" in textos
    assert "limitada a 30% dos honorários da fase 4" in textos and "termo de recebimento definitivo" in textos


def test_docx_em_fases_com_pacote_de_horas(cat, p, res):
    doc = _doc(cat, p, res, "B")
    linhas = [[c.text for c in row.cells] for row in doc.tables[4].rows]
    pacote = next(l for l in linhas if "pacote de 40 h/mês × 4 meses" in l[0])
    assert pacote[1] == "160 h"
    assert any(l[0].startswith("Hora excedente") and l[3] == "sob demanda" for l in linhas)
    textos = " ".join(_listas(doc))
    assert "o saldo não utilizado não acumula" in textos and "4 parcelas mensais iguais" in textos
    assert "Taxa de desempenho" not in textos
